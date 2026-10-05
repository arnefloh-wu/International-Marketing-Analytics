"""Build the deck "Introduction to marketing mix modelling" on the WU template (polish-slides skill).

Seven parts: what MMM is, how it works (with a key-terms recap), what we know (empirical
generalisations), doing it well, ad spend and attribution (last click, its problems, real-world tests and a
discussion), the reading "Analytics for Marketers" (Fantini & Narayandas, HBR 2023) with
an in-class discussion, and MMM in this course. Content and sources come from the resource search in
instructor/resources/mmm-resource-guide.md; every Alpenglow number comes from
make_mmm_figures.py (slides/figures/mmm_numbers.txt), which also draws the charts. Run it first.

usage: python slides/build_mmm_intro.py TEMPLATE.pptx OUT.pptx
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from wu_deck import (A, ACC, BODY, FIG, GAP, HEAD, INK, LIGHT, MONO, NAVY, NB, WHITE, Deck, I, box, card,
                     code_block, color, icon, ph, set_runs, stepper, text)

TEMPLATE, OUT = sys.argv[1:3]
FIGS = Path(__file__).resolve().parent / "figures"
FOOTER = "International Marketing Analytics" + NB + "·" + NB + "WT" + NB + "2026/27"
deck = Deck(TEMPLATE, FOOTER)
add_slide, L = deck.add_slide, deck.L
GREY_TXT = "595959"
PCT = NB + "%"


# ---------------------------------------------------------------- helpers ----
def divider(title, part, notes):
    s = add_slide("Kapitelfolie", title, notes=notes)
    set_runs(ph(s, 11).text_frame.paragraphs[0], part)
    return s


def badge(slide, x, y, num, name, d=0.34, fill=NAVY, size=12):
    b = box(slide, x, y, d, d, fill, name, shape=MSO_SHAPE.OVAL)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(num); r.font.size = Pt(size); r.font.bold = True
    color(r.font.color, WHITE)


def numbered(slide, x, y, w, items, name, row_h=0.62, gap=0.1, size=BODY):
    """Numbered rows: bg2 band, navy number badge, text with a bold lead-in."""
    for i, para in enumerate(items):
        ry = y + i * (row_h + gap)
        box(slide, x, ry, w, row_h, LIGHT, f"{name} row {i + 1}")
        badge(slide, x + 0.15, ry + (row_h - 0.34) / 2, i + 1, f"{name} number {i + 1}")
        text(slide, x + 0.65, ry, w - 0.8, row_h, [para], f"{name} text {i + 1}", size=size, anchor=MSO_ANCHOR.MIDDLE)


def callout(slide, y, h, icon_name, para, name, dark=False, x=0.5, w=9.0, size=BODY):
    box(slide, x, y, w, h, NAVY if dark else LIGHT, f"{name} box")
    icon(slide, icon_name, "white" if dark else "navy", x + 0.22, y + (h - 0.4) / 2, 0.4, name)
    if dark:
        lead, rest = para
        para = [(lead, {"bold": True, "col": WHITE}), (rest, {"col": WHITE})]
    text(slide, x + 0.82, y, w - 1.0, h, [para], f"{name} text", anchor=MSO_ANCHOR.MIDDLE, size=size)


def tile(slide, x, y, w, h, icon_name, head, lines, name, size=BODY):
    """Icon at the top left, bold navy heading, plain lines underneath."""
    box(slide, x, y, w, h, LIGHT, f"{name} tile")
    icon(slide, icon_name, "navy", x + 0.15, y + 0.15, 0.4, name)
    text(slide, x + 0.65, y + 0.15, w - 0.75, 0.4, [[(head, {"bold": True, "col": NAVY, "size": HEAD, "head": True})]],
         f"{name} heading", anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.15, y + 0.68, w - 0.3, h - 0.75, lines, f"{name} text", space=4, size=size)


def stat(slide, x, y, w, h, fig, label, name, fig_w=1.25, size=BODY):
    """Stat row: thin accent bar, 28 pt navy figure, explanation to its right."""
    box(slide, x, y, 0.07, h, ACC, f"{name} bar")
    text(slide, x + 0.2, y, fig_w, h, [[(fig, {"size": FIG, "bold": True, "col": NAVY, "head": True})]],
         f"{name} figure", anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.2 + fig_w, y, w - fig_w - 0.25, h, [label], f"{name} label", anchor=MSO_ANCHOR.MIDDLE, size=size)


def source(slide, txt, y=5.27, x=0.5, w=9.0):
    text(slide, x, y, w, 0.28, [[("Sources: " + txt, {"col": GREY_TXT, "size": 11})]], "Sources",
         anchor=MSO_ANCHOR.MIDDLE)


def picture(slide, fname, x, y, w, h, alt):
    pic = slide.shapes.add_picture(str(FIGS / fname), I(x), I(y), I(w), I(h))
    pic.name = "Chart " + fname
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def arrow(slide, x, y, name, w=0.22, h=0.36):
    box(slide, x, y, w, h, ACC, name, shape=MSO_SHAPE.CHEVRON)


def table(slide, rows, widths, name, size=BODY, head_size=HEAD, top=1.35, bold_first_col=True, row_margin=45000,
          mono_rows=()):
    gf = ph(slide, 1).insert_table(len(rows), len(rows[0]))
    gf.name = name
    tbl = gf.table
    for i, w in enumerate(widths):
        tbl.columns[i].width = I(w)
    gf.left, gf.top, gf.width = I(0.5), I(top), I(sum(widths))
    for r_i, row in enumerate(rows):
        for c_i, txt in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            p = cell.text_frame.paragraphs[0]
            set_runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(head_size if r_i == 0 else size)
                if c_i == 0 and r_i > 0 and bold_first_col:
                    r.font.bold = True
            cell.margin_top = cell.margin_bottom = Emu(row_margin)
    return tbl


def link(t, url):
    return (t, {"link": url, "col": ACC})


# ================================================================ slides ======
# 1  Title (same contact-block spacing as the other decks)
kontakt = next(sh for sh in L["Titelfolie Kontakt"].shapes if sh.name == "Kontaktdaten")
for para in kontakt.text_frame.paragraphs:
    if "Institute for International Business" in para.text:
        blank = copy.deepcopy(para._p)
        for r in blank.findall(f"{{{A}}}r"):
            blank.remove(r)
        end = etree.SubElement(blank, f"{{{A}}}endParaRPr"); end.set("lang", "en-GB"); end.set("sz", "1200")
        para._p.addnext(blank)
        break
s = add_slide("Titelfolie Kontakt", footer=None,
              notes="This lecture introduces marketing mix modelling, the method the whole course builds towards. No "
                    "prior analytics training is needed: we start from the business question and add the statistics step "
                    "by step. The content rests on a structured search of books, papers, industry reports and software "
                    "documentation; the sources are on the slides and on the further-reading slide at the end.")
set_runs(ph(s, 0).text_frame.paragraphs[0], "International Marketing Analytics")
set_runs(ph(s, 1).text_frame.paragraphs[0], "Introduction to marketing mix modelling")
set_runs(ph(s, 12).text_frame.paragraphs[0], "Session 1" + NB + "·" + NB + "WT" + NB + "2026/27")

# 2  The question
s = add_slide("Titel und Inhalt", "Where should the next euro go?",
              notes="Alpenglow is the course's fictional Vienna chocolate brand. In 2025 it spent about EUR 36 million on "
                    "media across six countries and five channels (51 weeks in the data; TV is about a third of it). "
                    "Every year the country managers ask for the same budget plus a few per cent, because nobody can say "
                    "what a euro in Polish paid social returns compared with a euro in French TV. That comparison is "
                    "exactly what marketing mix modelling (MMM) delivers. Numbers from slides/figures/mmm_numbers.txt.")
box(s, 0.5, 1.35, 4.3, 3.45, LIGHT, "Alpenglow panel")
text(s, 0.7, 1.43, 3.9, 0.4, [[("Alpenglow in 2025", {"bold": True, "col": NAVY, "size": HEAD, "head": True})]],
     "Alpenglow heading", anchor=MSO_ANCHOR.MIDDLE)
facts = [("6", "markets: AT, DE, FR, IT, NL, PL"), ("5", "channels: TV, online video, paid search, paid social, "
                                                       "out-of-home"),
         ("36" + NB + "m", "euros of media spend in 2025")]
for i, (fig, lab) in enumerate(facts):
    stat(s, 0.7, 1.95 + i * 0.93, 3.95, 0.8, fig, lab, f"Alpenglow fact {i + 1}", fig_w=1.3)
numbered(s, 4.95, 1.35, 4.55, [
    ("Contribution: ", "what did each channel add to sales?"),
    ("Return: ", "which euro earned more, TV in France or paid social in Poland?"),
    ("Allocation: ", "how should next year's fixed budget be split?")], "Questions", row_h=1.08,
    gap=0.105)
callout(s, 4.95, 0.6, "LuWallet", ("The CEO's brief: ", "the 2027 budget stays at the 2025 level. Move money only where "
                                                       "the next euro earns more."), "CEO", dark=True)

# ---------------------------------------------------------------- Part 1 -----
divider("What marketing mix modelling is", "Part 1 of 7" + NB + "·" + NB + "definition, inputs and outputs, "
        "MMM versus attribution and experiments, history",
        "First what MMM is and what it delivers, then how it relates to the other measurement tools, and why it is "
        "popular again.")

# 3  Definition
s = add_slide("Titel und Inhalt", "Marketing mix modelling in one sentence",
              notes="MMM is a statistical model of aggregate sales. It is old (the founding papers are from the 1970s) and "
                    "rests on the market response models of Hanssens, Parsons and Schultz. Little (1979) listed what such "
                    "a model must capture: carryover, diminishing returns, competition, interaction and dynamics. "
                    "Hanssens and Pauwels (2016) frame it as the link from marketing spend to financial value.")
box(s, 0.5, 1.35, 9.0, 1.15, NAVY, "Definition box")
icon(s, "LuLightbulb", "white", 0.72, 1.71, 0.42, "definition")
text(s, 1.35, 1.35, 7.95, 1.15, [[("MMM ", {"bold": True, "col": WHITE}),
                                  ("explains weekly sales with marketing and non-marketing drivers, splits sales into a "
                                   "base and the contribution of each activity, and turns these into budget decisions.",
                                   {"col": WHITE})]], "Definition text", anchor=MSO_ANCHOR.MIDDLE)
tw3 = (9.0 - GAP * 2) / 3
tiles = [("LuDatabase", "Aggregate data", ["Weekly totals per country and channel",
                                           "No tracking of individuals, so privacy-safe"]),
         ("LuSigma", "A regression model", ["Sales explained by spend, price, season and more",
                                            "Plus carryover and diminishing returns"]),
         ("LuWallet", "Decisions", ["What each euro returned", "Where the next euro should go"])]
for i, (ic, head, lines) in enumerate(tiles):
    tile(s, 0.5 + i * (tw3 + GAP), 2.65, tw3, 1.75, ic, head, [ln for ln in lines], f"Definition {i + 1}")
text(s, 0.5, 4.52, 2.1, 0.6, [[("Little (1979): ", {"bold": True, "col": NAVY}), ("a model must capture", {})]],
     "Little label", anchor=MSO_ANCHOR.MIDDLE, size=12)
chips = ["carryover", "diminishing returns", "competition", "interaction", "dynamics"]
chip_w = (9.0 - 2.2 - 0.1 * 4) / 5
for i, c in enumerate(chips):
    x = 2.7 + i * (chip_w + 0.1)
    box(s, x, 4.52, chip_w, 0.6, ACC, f"Little chip {i + 1}")
    text(s, x + 0.05, 4.52, chip_w - 0.1, 0.6, [[(c, {"bold": True, "col": WHITE})]], f"Little chip text {i + 1}",
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, size=12)
source(s, "Little 1979; Hanssens, Parsons & Schultz 2001; Hanssens & Pauwels 2016")

# 4  Inputs, model, outputs
s = add_slide("Titel und Inhalt", "Inputs, model, outputs",
              notes="Inputs: everything that moves sales, not only media. Leaving out price, distribution or season makes "
                    "the model credit media for their effect. Outputs: contributions (how many sales each driver added), "
                    "ROI or ROAS (return on ad spend: revenue per euro), marginal ROAS (revenue from the next euro, which "
                    "is what budget decisions need), response curves and an optimised budget split. In Alpenglow every "
                    "input is a column of data/mmm/alpenglow_weekly.csv.")
card(s, 0.5, 1.35, 3.5, 3.2, "Inputs", [
    ("Sales: ", "units or revenue per week"),
    ("Media: ", "spend by channel"),
    ("Marketing: ", "price, promotion, distribution"),
    ("Season: ", "trend, holidays, weather"),
    ("Context: ", "economy, competitors")], "LuDatabase", "Inputs card", head_h=0.45)
arrow(s, 4.08, 2.77, "Arrow in")
box(s, 4.38, 1.35, 1.24, 3.2, NAVY, "Model box")
icon(s, "LuSigma", "white", 4.77, 2.05, 0.46, "model")
text(s, 4.43, 2.6, 1.14, 1.3, [[("Model", {"bold": True, "col": WHITE, "size": HEAD, "head": True})],
                                [("regression with carryover and saturation", {"col": WHITE, "size": 11})]],
     "Model text", align=PP_ALIGN.CENTER, space=4)
arrow(s, 5.7, 2.77, "Arrow out")
card(s, 6.0, 1.35, 3.5, 3.2, "Outputs", [
    ("Contributions: ", "sales per driver"),
    ("ROAS: ", "revenue per euro spent"),
    ("Marginal ROAS: ", "revenue from the next euro"),
    ("Response curves: ", "sales at each spend level"),
    ("Allocation: ", "the best budget split")], "LuChartArea", "Outputs card", head_h=0.45)
callout(s, 4.75, 0.6, "LuFileSpreadsheet", [("In Alpenglow: ", {"bold": True, "col": NAVY}),
                                             ("every input is a column of ", {}),
                                             ("data/mmm/alpenglow_weekly.csv", {"font": MONO, "col": NAVY, "size": 12})],
        "Alpenglow inputs")

# 5  MMM versus attribution versus experiments
s = add_slide("Titel und Tabelle", "Three ways to measure marketing",
              notes="Multi-touch attribution (MTA) follows individual users' clicks and views and shares the credit for a "
                    "conversion among the touchpoints; it only sees digital channels and suffers from cookie loss. "
                    "Experiments (lift tests, geo experiments) switch spend on or off for a random group of users or "
                    "regions and are the most credible causal evidence, but cover one channel at a time. Google's Modern "
                    "Measurement Playbook calls the three a measurement tripod: use them together. Gordon et al. (2019) "
                    "show with 15 large Facebook experiments that observational methods often miss the true lift by a "
                    "wide margin.")
table(s, [
    ("", "Marketing mix modelling", "Multi-touch attribution", "Experiments"),
    ("Data", "weekly totals per channel and country", "clicks and views of individual users",
     "test and control groups or regions"),
    ("Answers", "what each channel added, offline too", "which touchpoints came before a sale",
     "the causal lift of one change"),
    ("Strength", "all channels; privacy-safe", "detailed and fast", "most credible evidence of cause"),
    ("Limit", "correlation; needs variation in spend", "digital only; credit is not cause",
     "one channel, one period; costly")],
    [1.35, 2.6, 2.55, 2.5], "Measurement table", size=13, row_margin=54000)
callout(s, 4.4, 0.75, "LuLayers", ("Triangulate: ", "experiments calibrate the MMM, the MMM covers every channel, "
                                                      "attribution steers digital campaigns day to day."), "Triangulate",
        dark=True)
source(s, "Google Modern Measurement Playbook 2023; IAB & MMA 2019; Gordon et al. 2019")

# 6  History and comeback
s = add_slide("Titel und Inhalt", "An old method, popular again",
              notes="MMM grew out of econometric market response models in the 1970s. Since the late 2010s, privacy rules "
                    "(GDPR), the loss of third-party cookies and Apple's App Tracking Transparency have made user-level "
                    "tracking less complete, so aggregate models that need no personal data are back. Google, Meta and "
                    "PyMC Labs released open-source MMM software. Gartner (2024) found that 64 % of senior marketing "
                    "leaders had adopted MMM; the WFA and Ebiquity (2026) report that 80 % of advertisers use MMM but "
                    "only 13 % turn the data into insight quickly.")
stops = [("1976", "Clarke: how long ad effects last"), ("1979", "Little: what a response model needs"),
         ("2001", "Hanssens et al.: Market Response Models"), ("2017", "Jin et al. (Google): Bayesian MMM"),
         ("2020s", "Open source: Robyn, PyMC-Marketing, Meridian")]
cg = 0.1
cw = (9.0 - cg * 4) / 5
road_y = 1.62
box(s, 0.5, road_y - 0.05, 8.75, 0.1, ACC, "Road")
end = box(s, 9.2, road_y - 0.15, 0.3, 0.3, ACC, "Road end", shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
end.rotation = 90
for i, (yr, desc) in enumerate(stops):
    x = 0.5 + i * (cw + cg)
    cx = x + cw / 2
    box(s, cx - 0.12, road_y - 0.12, 0.24, 0.24, NAVY, f"Milestone dot {i + 1}", shape=MSO_SHAPE.OVAL)
    box(s, x, 1.92, cw, 1.4, LIGHT, f"Milestone card {i + 1}")
    text(s, x + 0.1, 1.98, cw - 0.2, 0.42, [[(yr, {"bold": True, "col": NAVY, "size": 18, "head": True})]],
         f"Milestone year {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.1, 2.42, cw - 0.2, 0.86, [desc], f"Milestone text {i + 1}", size=12)
cw2 = (9.0 - GAP) / 2
card(s, 0.5, 3.5, cw2, 1.65, "Why now", [
    "GDPR, cookie loss and Apple's tracking rules weaken user-level tracking",
    "MMM needs no personal data"], "LuLock", "Why now card", head_h=0.45)
box(s, 0.5 + cw2 + GAP, 3.5, cw2, 1.65, LIGHT, "Adoption panel")
stat(s, 0.65 + cw2 + GAP, 3.6, cw2 - 0.25, 0.68, "64" + PCT, "of senior marketing leaders use MMM (Gartner 2024)",
     "Adoption 1", fig_w=1.35, size=12)
stat(s, 0.65 + cw2 + GAP, 4.37, cw2 - 0.25, 0.68, "13" + PCT, "of advertisers turn it into insight quickly "
                                                            "(WFA 2026)", "Adoption 2", fig_w=1.35, size=12)
source(s, "Clarke 1976; Little 1979; Hanssens et al. 2001; Jin et al. 2017; Gartner 2024; WFA & Ebiquity 2026")

# ---------------------------------------------------------------- Part 2 -----
divider("How MMM works", "Part 2 of 7" + NB + "·" + NB + "base and incremental sales, the equation, adstock, "
        "saturation, seasonality",
        "Five building blocks. Each one is a variable or a transformation in the regression you will estimate.")

# 7  Base plus incremental
s = add_slide("Titel und Inhalt", "Sales = base + incremental",
              notes="The chart splits German weekly sales in 2025 into a base and the part the model attributes to media. "
                    "It comes from a deliberately simple regression (all five channels with the same carryover of 0.5, no "
                    "saturation); in this first model media account for about 41 % of fitted sales. That number is too "
                    "high: Alpenglow raises its budgets before Christmas, when people buy chocolate anyway, so a naive "
                    "model credits media with part of the seasonal demand (the endogeneity pitfall later in this deck). "
                    "Session 2 builds a proper model and you will see the media share fall considerably. The dashed "
                    "line is actual sales; the gap to the coloured area is what the model cannot explain.")
picture(s, "mmm_decomposition.png", 0.5, 1.35, 5.4, 3.5,
        "Stacked area chart of weekly sales in Germany in 2025, split into a base of about 300 to 800 thousand units and "
        "a media part on top; the dashed line of actual sales follows the total closely, with a Christmas peak")
rows = [("Base: ", "what would sell without media: price, distribution, season, trend"),
        ("Incremental: ", "what media add on top, channel by channel"),
        ("Contribution: ", "the sales that would disappear if a channel were switched off")]
for i, para in enumerate(rows):
    y = 1.35 + i * 1.0
    box(s, 6.05, y, 3.45, 0.9, LIGHT, f"Decomp row {i + 1}")
    box(s, 6.05, y, 0.07, 0.9, ACC, f"Decomp bar {i + 1}")
    text(s, 6.25, y, 3.17, 0.9, [para], f"Decomp text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
box(s, 6.05, 4.35, 3.45, 0.5, NAVY, "Decomp share box")
text(s, 6.2, 4.35, 3.25, 0.5, [[("Naive model: ", {"bold": True, "col": WHITE, "size": 12}),
                                ("media ≈ 41" + PCT + ", too high", {"col": WHITE, "size": 12})]], "Decomp share text",
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.5, 4.95, 9.0, 0.32, [[("Alpenglow Germany, 2025, naive regression: budgets follow Christmas demand, so media get credit for "
                                 "seasonal sales (see the pitfalls slide).", {"col": GREY_TXT, "size": 11})]], "Decomp caption", anchor=MSO_ANCHOR.MIDDLE)
source(s, "Hanssens et al. 2001; Jin et al. 2017; own calculation on the Alpenglow data")

# 8  The equation
s = add_slide("Titel und Inhalt", "The regression behind it",
              notes="Each channel's weekly spend is first transformed (carryover, then diminishing returns) and multiplied "
                    "by a coefficient; the controls form the base. In the additive form a coefficient is extra units per "
                    "unit of transformed spend. In the log-log form coefficients are elasticities: the per cent change in "
                    "sales for a 1 % change in the driver. Log-log is the classic constant-elasticity model of Hanssens et "
                    "al.; Tellis (2006) gives a compact primer. In Session 2 you estimate both with statsmodels.")
stepper(s, 0.5, 1.35, 9.0, 0.66, ["Spend per\vchannel", "Adstock:\vcarryover", "Saturation:\vless per €",
                                  "× β:\vthe effect", "Sales\vcontribution"], "Pipeline", size=11)
cw2 = (9.0 - GAP) / 2
forms = [("Additive form", "sales_t = β0 + Σ βc·f(spend_c,t)\n          + γ·controls_t + ε_t",
          [("βc: ", "extra units per unit of transformed spend"), ("Good for: ", "contributions in units or euros")]),
         ("Log-log form", "ln sales_t = β0 + βp·ln price_t\n   + Σ βc·ln(1 + adstock_c,t) + … + ε_t",
          [("β: ", "elasticity, % change in sales per 1" + PCT + " change"), ("Good for: ", "comparing countries and drivers")])]
for k, (head, eq, lines) in enumerate(forms):
    x = 0.5 + k * (cw2 + GAP)
    box(s, x, 2.2, cw2, 0.45, ACC, f"Form {k + 1} header")
    text(s, x + 0.15, 2.2, cw2 - 0.3, 0.45, [head], f"Form {k + 1} heading", size=HEAD, col=WHITE, bold=True, head=True,
         anchor=MSO_ANCHOR.MIDDLE)
    code_block(s, x, 2.65, cw2, 0.85, eq, f"Form {k + 1} equation", size=11)
    box(s, x, 3.5, cw2, 1.0, LIGHT, f"Form {k + 1} body")
    text(s, x + 0.18, 3.58, cw2 - 0.33, 0.9, lines, f"Form {k + 1} text", size=12, bullets=True, space=4)
callout(s, 4.65, 0.6, "LuSlidersHorizontal", ("Controls: ", "price, promotion, distribution, trend, season, holidays, "
                                                           "weather, economy, competitor spend."), "Controls", size=12)
source(s, "Hanssens et al. 2001; Tellis 2006; Jin et al. 2017")

# 9  Adstock
s = add_slide("Titel und Inhalt", "Adstock: advertising keeps working",
              notes="Geometric adstock carries a share lambda (the decay) of last week's advertising pressure into this "
                    "week. The chart shows one burst of EUR 100k: with decay 0.3 the effect is almost gone after three "
                    "weeks; with 0.7 it lasts about two months. Half-life is the number of weeks until half the pressure is "
                    "gone: ln 0.5 / ln lambda. The total pressure is 1/(1 - lambda) times the first week: 1.4 versus 3.3. "
                    "Jin et al. (2017) also use a delayed adstock with a peak after some weeks; Robyn offers Weibull "
                    "variants. Synthetic example, made in make_mmm_figures.py.")
picture(s, "mmm_adstock.png", 0.5, 1.35, 5.2, 3.3,
        "Line chart: a spend of 100 in week 1 decays to 30 in week 2 and almost zero by week 5 with decay 0.3, and to 70, 49, "
        "34 and so on with decay 0.7, still about 2 in week 12")
code_block(s, 5.9, 1.35, 3.6, 0.85, "A_t = spend_t + λ·A_(t−1)", "Adstock formula", size=12,
           caption="Geometric adstock")
rows = [("λ (decay): ", "share of last week's pressure that carries over"),
        ("Half-life: ", "ln" + NB + "0.5" + NB + "/" + NB + "ln" + NB + "λ = 0.6 or 1.9 weeks here"),
        ("Total effect: ", "1/(1−λ) × week 1: 1.4× versus 3.3×")]
for i, para in enumerate(rows):
    y = 2.35 + i * 0.78
    box(s, 5.9, y, 3.6, 0.7, LIGHT, f"Adstock row {i + 1}")
    box(s, 5.9, y, 0.07, 0.7, ACC, f"Adstock bar {i + 1}")
    text(s, 6.1, y, 3.33, 0.7, [para], f"Adstock text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
callout(s, 4.8, 0.45, "LuHourglass", ("Without adstock, ", "the weeks after a campaign are credited to whatever else "
                                                          "happened then."), "Adstock why", dark=True, size=12)
source(s, "Little 1979; Jin et al. 2017; Robyn documentation (geometric and Weibull adstock)")

# 10  Saturation
s = add_slide("Titel und Inhalt", "Saturation: each extra euro buys less",
              notes="The Hill curve rises slowly, then steeply, then flattens. K is the spend at which half the maximum "
                    "effect is reached; the slope s sets the shape: s of 1 or less bends from the start (C-shape), s above "
                    "1 gives an S-shape. In this synthetic example the next EUR 10k buys about 13,000 extra units at EUR 30k "
                    "a week, but fewer than 2,000 at EUR 150k. Average ROAS can look good while marginal ROAS is poor: "
                    "budgets should move to where the next euro earns most. Hill is the default in Meridian and Robyn; a "
                    "logistic curve is an alternative. Adstock and saturation are hard to separate in the data (Meridian "
                    "documentation).")
picture(s, "mmm_saturation.png", 0.5, 1.35, 5.2, 3.3,
        "S-shaped curve of extra sales against weekly spend from 0 to 200 thousand euros, flattening towards 120 thousand "
        "units; the next 10 thousand euros add 13 thousand units at 30 thousand euros spend but only 1.8 thousand at 150")
code_block(s, 5.9, 1.35, 3.6, 0.85, "Hill(x) = x^s / (x^s + K^s)", "Hill formula", size=12, caption="Hill saturation")
rows = [("K: ", "spend at half the maximum effect"),
        ("s: ", "shape; above 1 gives an S-curve"),
        ("Marginal ROAS: ", "revenue from the next euro, the slope of the curve")]
for i, para in enumerate(rows):
    y = 2.35 + i * 0.78
    box(s, 5.9, y, 3.6, 0.7, LIGHT, f"Hill row {i + 1}")
    box(s, 5.9, y, 0.07, 0.7, ACC, f"Hill bar {i + 1}")
    text(s, 6.1, y, 3.33, 0.7, [para], f"Hill text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
callout(s, 4.8, 0.45, "LuGauge", ("Decide at the margin: ", "move money to where the next euro earns most."),
        "Hill why", dark=True, size=12)
source(s, "Jin et al. 2017; Google Meridian documentation (media saturation and lagging)")

# 11  Seasonality, trend and holidays
s = add_slide("Titel und Inhalt", "Seasonality, trend and holidays",
              notes="German weekly sales peak before Christmas, Easter and Valentine's Day and dip in summer (chocolate and "
                    "heat do not mix). TV spend peaks at the same time: in Christmas weeks it averages about EUR 188k a week, "
                    "against about EUR 100k in other weeks. Without controls for season and holidays, the model would give "
                    "TV the credit for Christmas. Seasonality is usually modelled with a few Fourier terms (pairs of sine "
                    "and cosine waves), as in Prophet, PyMC-Marketing and Meridian; Hyndman and Athanasopoulos explain how "
                    "many to use.")
picture(s, "mmm_de_sales_tv.png", 0.5, 1.35, 5.6, 3.4,
        "Two panels for Germany 2023 to 2025: weekly sales between about 450 and 1,750 thousand units with peaks before "
        "each Christmas, and weekly TV spend with bursts of up to about 480 thousand euros in spring and before Christmas; "
        "Christmas weeks are shaded grey")
text(s, 0.5, 4.75, 5.6, 0.3, [[("Alpenglow Germany, 2023–2025; grey: Christmas selling weeks.",
                                {"col": GREY_TXT, "size": 11})]], "Season caption", anchor=MSO_ANCHOR.MIDDLE)
rows = [("Trend: ", "slow growth or decline, a week counter"),
        ("Seasonality: ", "the yearly wave, Fourier terms"),
        ("Holidays: ", "dummies for Christmas, Easter, Valentine's"),
        ("Weather, economy: ", "temperature, consumer confidence")]
for i, para in enumerate(rows):
    y = 1.35 + i * 0.86
    box(s, 6.25, y, 3.25, 0.78, LIGHT, f"Season row {i + 1}")
    box(s, 6.25, y, 0.07, 0.78, ACC, f"Season bar {i + 1}")
    text(s, 6.45, y, 2.98, 0.78, [para], f"Season text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
callout(s, 5.1, 0.45, "LuCalendarDays", ("TV peaks when sales peak: ", "does TV drive Christmas, or Christmas the TV "
                                                                      "budget?"), "Season why", dark=True, size=12)

# 11b  Key terms: the vocabulary of MMM (two slides, six terms each)
TERMS = [
    ("Key terms (1): what the model splits",
     "Six words you will use every week. Each card has the definition and where you meet it in the Alpenglow data.",
     [("LuSlidersHorizontal", "Marketing mix", "The levers a brand controls: product, price, place (distribution) and promotion.",
       "Alpenglow: price, promotions, distribution and five media channels"),
      ("LuHouse", "Base sales", "What would sell without media: driven by price, distribution, season and trend.",
       "The dark area in the decomposition chart"),
      ("LuTrendingUp", "Incremental sales", "What a channel adds on top of the base: the sales lost if it were switched off.",
       "Also called the channel's contribution"),
      ("LuHistory", "Adstock", "Advertising keeps working after the week it runs; the decay is the share that carries over.",
       "Decay 0.7: half the effect is gone after about 2 weeks"),
      ("LuGauge", "Saturation", "Each extra euro buys less: the response flattens as spend grows.",
       "The next EUR 10k buys far fewer units at high spend"),
      ("LuShieldCheck", "Control variables", "Non-media drivers in the model, so media are not credited with their effect.",
       "Price, promotions, holidays, temperature, competitors")]),
    ("Key terms (2): what the model tells you",
     "These six words turn model output into decisions. Marginal ROAS is the one that moves budgets.",
     [("LuScale", "Elasticity", "The % change in sales for a 1 % change in a driver; the coefficient in a log-log model.",
       "Price elasticity −1.4: 1 % higher price, 1.4 % fewer units"),
      ("LuChartSpline", "Response curve", "Incremental sales at each spend level: adstock and saturation in one picture.",
       "Read it to see where a channel stops paying off"),
      ("LuCoins", "ROAS and ROI", "ROAS: incremental revenue per euro spent. ROI: incremental profit per euro spent.",
       "ROAS 3: every euro returned EUR 3 of revenue on average"),
      ("LuTarget", "Marginal ROAS", "The return on the next euro: the slope of the response curve.",
       "Move budget to the highest marginal, not average, ROAS"),
      ("LuCalendarDays", "Holdout", "Fit on past weeks, predict weeks the model has not seen, compare the errors.",
       "Fit 2023–2024, forecast 2025, report the MAPE"),
      ("LuFlaskConical", "Calibration", "Checking or anchoring model effects with experiments such as geo lift tests.",
       "Paid-social test in 12 of 40 German regions")]),
]
tw3 = (9.0 - GAP * 2) / 3
th = (4.25 - GAP) / 2
for title, notes, terms in TERMS:
    s = add_slide("Titel und Inhalt", title, notes=notes)
    for i, (ic, term, definition, example) in enumerate(terms):
        x = 0.5 + (i % 3) * (tw3 + GAP)
        y = 1.35 + (i // 3) * (th + GAP)
        tile(s, x, y, tw3, th, ic, term, [definition, [(example, {"col": ACC, "bold": True, "size": 11.5})]],
             f"Term {i + 1}", size=12)

# ---------------------------------------------------------------- Part 3 -----
divider("What we know: empirical generalisations", "Part 3 of 7" + NB + "·" + NB + "elasticities, carryover and "
        "wear-out, differences between countries",
        "Decades of studies give benchmarks. Use them to judge whether your own model's results are plausible and, in a "
        "Bayesian model, as priors.")

# 12  Elasticities
s = add_slide("Titel und Inhalt", "Elasticities: price versus advertising",
              notes="An elasticity is the per cent change in sales for a 1 % change in a driver. Sethuraman, Tellis and "
                    "Briesch (2011) pooled published brand advertising elasticities: on average 0.12 in the short term and "
                    "0.24 in the long term. Henningsen, Heuke and Clement (2011) find 0.09 in an international database. "
                    "Lodish et al. (1995): about half of 389 TV weight tests showed no significant sales effect, and "
                    "Shapiro, Hitsch and Tuchman (2021) find small TV elasticities and negative marginal ROI for most of 288 "
                    "brands. On the right: a first, simple log-log regression on Alpenglow Germany gives a price elasticity "
                    "of about -1.4 and a TV elasticity of about 0.06. Price moves sales far more per per cent than "
                    "advertising does; Session 2 estimates these properly, by country.")
box(s, 0.5, 1.35, 4.5, 3.55, LIGHT, "Literature panel")
text(s, 0.7, 1.42, 4.1, 0.4, [[("Advertising: the meta-analyses", {"bold": True, "col": NAVY, "size": HEAD,
                                                                   "head": True})]], "Literature heading",
     anchor=MSO_ANCHOR.MIDDLE)
lit = [("0.12", "short-term elasticity, mean (Sethuraman et al. 2011)"),
       ("0.24", "long-term elasticity, mean (same study)"),
       ("0.09", "mean in an international database (Henningsen et al. 2011)"),
       ("≈" + NB + "½", "of 389 TV weight tests: no significant sales effect (Lodish et al. 1995)")]
for i, (fig, lab) in enumerate(lit):
    stat(s, 0.7, 1.9 + i * 0.74, 4.15, 0.64, fig, lab, f"Literature stat {i + 1}", fig_w=1.05, size=12)
card_x, card_w = 5.15, 4.35
box(s, card_x, 1.35, card_w, 0.45, ACC, "Alpenglow header")
text(s, card_x + 0.15, 1.35, card_w - 0.75, 0.45, ["Alpenglow Germany: first look"], "Alpenglow heading", size=HEAD,
     col=WHITE, bold=True, head=True, anchor=MSO_ANCHOR.MIDDLE)
icon(s, "LuCandy", "white", card_x + card_w - 0.48, 1.4, 0.34, "Alpenglow")
box(s, card_x, 1.8, card_w, 3.1, LIGHT, "Alpenglow body")
stat(s, card_x + 0.2, 1.95, card_w - 0.35, 0.85, "−1.4", "price: 1" + PCT + " lower price, about 1.4" + PCT + " more units",
     "Price elasticity", fig_w=1.15, size=12)
stat(s, card_x + 0.2, 2.95, card_w - 0.35, 0.85, "0.06", "TV: 1" + PCT + " more TV spend, about 0.06" + PCT +
     " more units", "TV elasticity", fig_w=1.15, size=12)
text(s, card_x + 0.2, 3.95, card_w - 0.35, 0.85, [[("Simple log-log regression, 2023–2025. You estimate it properly, "
                                                     "by country, in Session 2.", {"col": GREY_TXT, "size": 11})]],
     "Alpenglow note", anchor=MSO_ANCHOR.MIDDLE)
callout(s, 5.02, 0.53, "LuScale", ("Reading 0.12: ", "10" + PCT + " more advertising lifts sales by about 1.2" + PCT + "."),
        "Elasticity reading", dark=True, size=12)

# 13  Carryover and wear-out
s = add_slide("Titel und Inhalt", "Carryover and wear-out",
              notes="Profit Ability 2 (Ebiquity for Thinkbox, 2024) pooled UK models for 141 brands: about 60 % of the "
                    "profit effect of advertising arrived after the first 13 weeks. The long-term elasticity in Sethuraman "
                    "et al. is about twice the short-term one. Clarke (1976) showed that estimated carryover depends on "
                    "the data interval: annual or monthly data imply much longer effects than weekly data, a bias that "
                    "Köhler et al. (2017) confirm in their meta-analysis of carryover. Wear-in and wear-out (Tellis 2004; "
                    "Vakratsas and Ambler 1999): an ad may need repetition before it works and loses effect when shown "
                    "too often or too long.")
box(s, 0.5, 1.35, 4.3, 3.4, LIGHT, "Carryover panel")
text(s, 0.7, 1.42, 3.9, 0.4, [[("How long effects last", {"bold": True, "col": NAVY, "size": HEAD, "head": True})]],
     "Carryover heading", anchor=MSO_ANCHOR.MIDDLE)
stat(s, 0.7, 1.95, 3.95, 1.25, "60" + PCT, "of the effect came after week 13 (141 UK brands; Ebiquity & Thinkbox "
                                           "2024)", "Carryover stat 1", fig_w=1.3, size=12)
stat(s, 0.7, 3.35, 3.95, 1.25, "2×", "long-term elasticity versus short-term: 0.24 against 0.12 (Sethuraman et al. "
                                     "2011)", "Carryover stat 2", fig_w=1.3, size=12)
tile(s, 4.95, 1.35, 4.55, 1.6, "LuCalendarDays", "Data interval matters",
     ["Monthly or annual data suggest longer carryover than weekly data (Clarke 1976; Köhler et al. 2017)"],
     "Interval", size=12)
tile(s, 4.95, 3.12, 4.55, 1.63, "LuRepeat", "Wear-in and wear-out",
     ["Ads may need repetition to work, and lose effect when shown too often or too long (Tellis 2004; Vakratsas & "
      "Ambler 1999)"], "Wearout", size=12)
callout(s, 4.92, 0.63, "LuHourglass", ("For your model: ", "estimate carryover per channel from weekly data and check "
                                                         "half-lives against these benchmarks."), "Carryover rule",
        dark=True, size=12)

# 14  International differences
s = add_slide("Titel und Inhalt", "Effects differ between countries",
              notes="Bahadir, Bharadwaj and Srivastava (2015) model 14 emerging and developed markets: advertising and "
                    "product innovation matter more in emerging markets, while price and distribution effects differ by "
                    "country type. Deleersnyder et al. (2009): across 37 countries and 25 years advertising budgets rise and "
                    "fall with the economy, and how strongly depends on national culture. Van Heerde et al. (2013): in "
                    "downturns price sensitivity rises and advertising effectiveness falls. Fischer et al. (2011) used "
                    "response elasticities to reallocate Bayer's budget across countries, products and activities and "
                    "report a profit improvement of about EUR 500 million. For Alpenglow, data/mmm/country_meta.csv holds "
                    "GDP per capita, Hofstede scores and retail concentration to explain such differences.")
intl = [("LuEarth", "Market type: ", "advertising matters more in emerging markets; price and distribution effects "
                                     "differ too (Bahadir et al. 2015, 14 markets)"),
        ("LuTrendingDown", "Business cycle: ", "in downturns price sensitivity rises and advertising works less "
                                               "(van Heerde et al. 2013)"),
        ("LuUsers", "Culture: ", "how strongly ad budgets follow the economy depends on national culture "
                                 "(Deleersnyder et al. 2009, 37 countries)"),
        ("LuCoins", "Payoff: ", "reallocating across countries with elasticities: about EUR" + NB + "500m more "
                                "profit at Bayer (Fischer et al. 2011)")]
for i, (ic, lead, rest) in enumerate(intl):
    y = 1.35 + i * 0.84
    box(s, 0.5, y, 9.0, 0.76, LIGHT, f"Country row {i + 1}")
    box(s, 0.5, y, 0.07, 0.76, ACC, f"Country bar {i + 1}")
    icon(s, ic, "navy", 0.75, y + 0.17, 0.42, f"country {i + 1}")
    text(s, 1.35, y, 8.0, 0.76, [(lead, rest)], f"Country text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
callout(s, 4.75, 0.8, "LuMap", ("For Alpenglow: ", "expect Poland to respond differently from Germany. Hierarchical "
                                                   "models let countries differ while borrowing strength from each other "
                                                   "(Sun et al. 2017)."), "Country rule", dark=True, size=12)

# ---------------------------------------------------------------- Part 4 -----
divider("Doing it well", "Part 4 of 7" + NB + "·" + NB + "data, pitfalls, validation, Bayesian priors, open-source tools",
        "Most MMM failures are data and identification problems, not software problems.")

# 15  Data requirements
s = add_slide("Titel und Inhalt", "What the data must offer",
              notes="Two to three years of weekly data is the usual starting point: the seasonal pattern must repeat "
                    "before a model can separate it from media. Alpenglow has 156 weeks for each of six countries. Spend "
                    "must vary: a channel that runs at the same level every week, or always together with another, cannot "
                    "be measured. Nielsen's best-practice guide (with Google and Meta, 2022) recommends modelling at the "
                    "most granular geography available and including non-marketing drivers: in its comparison a "
                    "media-only model overstated advertising ROI by 68 %.")
tw2 = (9.0 - GAP) / 2
tiles4 = [("LuCalendarDays", "Two to three years, weekly", ["Seasons must repeat; Alpenglow has 156 weeks per country"]),
          ("LuChartColumn", "Variation in spend", ["Channels must move up and down, and not always together"]),
          ("LuMap", "Fine geography", ["Countries or regions, not one national total (Nielsen 2022)"]),
          ("LuListChecks", "All drivers", ["Price, distribution, season and competitors, not only media"])]
for i, (ic, head, lines) in enumerate(tiles4):
    x = 0.5 + (i % 2) * (tw2 + GAP)
    y = 1.35 + (i // 2) * (1.45 + GAP)
    tile(s, x, y, tw2, 1.45, ic, head, lines, f"Data tile {i + 1}")
box(s, 0.5, 4.5, 9.0, 0.95, NAVY, "Media-only box")
text(s, 0.75, 4.5, 1.6, 0.95, [[("+68" + PCT, {"size": FIG, "bold": True, "col": WHITE, "head": True})]],
     "Media-only figure", anchor=MSO_ANCHOR.MIDDLE)
text(s, 2.4, 4.5, 6.9, 0.95, [[("A media-only model ", {"bold": True, "col": WHITE}),
                               ("overstated advertising ROI by 68" + PCT + " compared with a model that included the "
                                "other drivers (Nielsen 2022).", {"col": WHITE})]], "Media-only text",
     anchor=MSO_ANCHOR.MIDDLE, size=12)

# 16  Pitfalls
s = add_slide("Titel und Inhalt", "Four pitfalls",
              notes="Chan and Perry (2017) list why MMM is hard. Endogeneity: budgets follow demand, so spend and sales "
                    "rise together even without any effect; in Alpenglow Germany TV spend almost doubles in Christmas weeks. "
                    "Multicollinearity: channels launched together cannot be separated; check variance inflation factors "
                    "(VIF). Overfitting: five channels with a decay, a shape and an effect each are many parameters for 156 "
                    "weeks. Identification: saturation and effects that change over time can mimic each other (Dew, "
                    "Padilla and Shchetkina 2024). Heusch (2026, Zalando) shows a standard MMM reporting a paid-search "
                    "ROAS of 10.6 where geo experiments found 4.2.")
pit = [("LuRepeat", "Endogeneity", ["Budgets follow demand: German TV spend nearly doubles in Christmas weeks"]),
       ("LuShuffle", "Multicollinearity", ["Channels that move together cannot be told apart; check the VIF"]),
       ("LuSlidersHorizontal", "Overfitting", ["Many parameters, few weeks: the fit looks great, the forecast fails"]),
       ("LuCrosshair", "Identification", ["Saturation and effects that change over time can mimic each other"])]
for i, (ic, head, lines) in enumerate(pit):
    x = 0.5 + (i % 2) * (tw2 + GAP)
    y = 1.35 + (i // 2) * (1.45 + GAP)
    card(s, x, y, tw2, 1.45, head, lines, ic, f"Pitfall {i + 1}", head_h=0.45, bullets=False)
box(s, 0.5, 4.5, 9.0, 0.95, NAVY, "Heusch box")
text(s, 0.75, 4.5, 1.9, 0.95, [[("10.6 vs 4.2", {"size": 22, "bold": True, "col": WHITE, "head": True})]],
     "Heusch figure", anchor=MSO_ANCHOR.MIDDLE)
text(s, 2.75, 4.5, 6.55, 0.95, [[("Paid-search ROAS: ", {"bold": True, "col": WHITE}),
                                 ("a standard MMM against geo experiments at Zalando (Heusch 2026).", {"col": WHITE})]],
     "Heusch text", anchor=MSO_ANCHOR.MIDDLE, size=12)
source(s, "Chan & Perry 2017; Dew et al. 2024; Heusch 2026", y=5.47)

# 17  Validation
s = add_slide("Titel und Inhalt", "Validate before you trust",
              notes="A good in-sample fit says little. First, hold back the last weeks and forecast them. Second, check "
                    "plausibility: positive media effects, half-lives and elasticities in the ranges of Part 3. Third, "
                    "calibrate with experiments: geo experiments switch spend up or down in some regions and compare them "
                    "with control regions (Vaver and Koehler 2011). Gordon et al. (2019) showed with 15 large Facebook "
                    "experiments, and Gordon, Moakler and Zettelmeyer (2023) with 663, that non-experimental estimates are "
                    "often far off. Alpenglow ran such a test: 2.5 times paid-social spend in 12 of 40 German regions for "
                    "8 weeks (data/experiments/geolift_germany.csv).")
stepper(s, 0.5, 1.35, 9.0, 0.6, ["Fit", "Holdout\vforecast", "Plausibility\vchecks", "Calibrate with\vexperiments"],
        "Validation steps", size=12, current=3)
tw3 = (9.0 - GAP * 2) / 3
val = [("LuChartLine", "Holdout", ["Fit to 2023–2024", "Forecast 2025", "Compare errors (MAPE)"]),
       ("LuListChecks", "Plausibility", ["Media effects positive", "Half-lives and elasticities near benchmarks",
                                         "Stable when data change"]),
       ("LuFlaskConical", "Experiments", ["Spend up or off in test regions", "Lift = test minus control",
                                          "Alpenglow: paid social, 12 of 40 German regions"])]
for i, (ic, head, items) in enumerate(val):
    card(s, 0.5 + i * (tw3 + GAP), 2.15, tw3, 2.25, head, items, ic, f"Validation card {i + 1}", head_h=0.45)
callout(s, 4.55, 0.85, "LuTriangleAlert", ("Why experiments: ", "across 15 and then 663 Facebook experiments, "
                                                              "non-experimental estimates were often far from the true "
                                                              "lift."), "Gordon", dark=True, size=12)
source(s, "Vaver & Koehler 2011; Gordon et al. 2019; Gordon, Moakler & Zettelmeyer 2023", y=5.43)

# 18  Bayesian MMM
s = add_slide("Titel und Inhalt", "Bayesian MMM: start from what we know",
              notes="Bayesian estimation combines a prior (what we believe before seeing the data) with the data and "
                    "returns a posterior: a full distribution for every effect, so each ROAS comes with a range. Priors "
                    "can come from empirical generalisations such as Sethuraman et al. (Hanssens 2018 argues for this use) "
                    "or from lift tests: Google's calibration paper (Zhang et al. 2024) puts experiment-based priors on "
                    "each channel's ROI, the approach Meridian, Robyn and PyMC-Marketing now support. Jin et al. (2017) is "
                    "the reference model. The risk: with weak data the prior decides, so compare prior and posterior.")
flow = [("LuBookOpen", "Prior", "what we believe first: elasticities near 0.1, a lift-test result"),
        ("LuDatabase", "Data", "156 weeks of sales, spend and controls per country"),
        ("LuChartArea", "Posterior", "updated belief: every effect and ROAS with a range")]
fw = 2.7
for i, (ic, head, desc) in enumerate(flow):
    x = 0.5 + i * (fw + 0.45)
    fill = NAVY if i == 2 else LIGHT
    box(s, x, 1.35, fw, 1.45, fill, f"Bayes box {i + 1}")
    icon(s, ic, "white" if i == 2 else "navy", x + 0.15, 1.47, 0.4, f"bayes {i + 1}")
    tc = WHITE if i == 2 else NAVY
    text(s, x + 0.65, 1.47, fw - 0.75, 0.4, [[(head, {"bold": True, "col": tc, "size": HEAD, "head": True})]],
         f"Bayes head {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, 1.95, fw - 0.3, 0.8, [[(desc, {"col": WHITE if i == 2 else INK, "size": 12})]],
         f"Bayes text {i + 1}")
    if i < 2:
        text(s, x + fw, 1.35, 0.45, 1.45, [[("+" if i == 0 else "→", {"bold": True, "col": ACC, "size": 26})]],
             f"Bayes operator {i + 1}", anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
cw2 = (9.0 - GAP) / 2
card(s, 0.5, 3.0, cw2, 2.1, "Why it suits MMM", [
    "Short series, many parameters: priors keep estimates sensible",
    "Uncertainty for every output",
    "Experiments enter as priors"], "LuSparkles", "Bayes why card", head_h=0.45)
card(s, 0.5 + cw2 + GAP, 3.0, cw2, 2.1, "What to watch", [
    "Weak data: the prior decides; compare prior and posterior",
    "Check that the sampler converged",
    "State and justify every prior"], "LuEye", "Bayes watch card", head_h=0.45)
source(s, "Jin et al. 2017; Hanssens 2018; Zhang et al. 2024 (Google)")

# 19  Open-source tools
s = add_slide("Titel und Tabelle", "Open-source MMM tools",
              notes="All three are free. Robyn (Meta) fits ridge regressions and searches over adstock and saturation "
                    "settings with an evolutionary algorithm (Nevergrad), then picks models on fit and plausibility; its "
                    "Python port is still a beta. Meridian (Google) is a hierarchical Bayesian model across geos with reach "
                    "and frequency and ROI priors; it replaced LightweightMMM, which Google archived in January 2026. "
                    "PyMC-Marketing (PyMC Labs) offers adstock and saturation building blocks, lift-test calibration, "
                    "time-varying effects, geo models and a budget optimiser. Version numbers as listed on PyPI and GitHub "
                    "on 5 October 2026. A 2026 open-access review in Customer Needs and Solutions compares the three.")
table(s, [
    ("", "Robyn", "Meridian", "PyMC-Marketing"),
    ("Maker", "Meta", "Google", "PyMC Labs"),
    ("Language", "R (Python beta)", "Python", "Python"),
    ("Method", "ridge regression with automated search", "hierarchical Bayesian, by geo", "Bayesian, modular"),
    ("Strengths", "calibration, budget allocator, widely used", "many geos, reach and frequency, ROI priors",
     "lift-test calibration, time-varying effects, optimiser"),
    ("Version", "3.12.0 (Dec 2024)", "2.1.0 (Sep 2026)", "1.2.0 (Sep 2026)")],
    [1.5, 2.5, 2.5, 2.5], "Tools table", size=13, row_margin=60000)
callout(s, 4.62, 0.6, "LuCode", ("In this course: ", "PyMC-Marketing and Meridian in Python, Session 5."), "Tools course",
        size=12)
source(s, "Runge et al. 2024; Meridian and PyMC-Marketing documentation; open-source MMM review (2026)")

# ---------------------------------------------------------------- Part 5 -----
divider("Ad spend and attribution", "Part 5 of 7" + NB + "·" + NB + "where the money goes, last click, why "
        "attribution misleads, real-world tests", "Before we read the article: how much money is at stake, and how most "
        "firms decide today which ad deserves the credit for a sale.")

# M1  Ad spend
s = add_slide("Titel und Inhalt", "Where the advertising money goes",
              notes="WPP Media's This Year Next Year (mid-year 2025) puts global advertising revenue at about $1.14 trillion "
                    "in 2025, up 8.8 %, with digital at about 73 % (search, social, online video, retail media); its "
                    "forecasts for 2026 are around $1.3 trillion. In Europe, IAB Europe's AdEx Benchmark 2024 counts "
                    "€118.9 billion of digital advertising, up 16 %, 67 % of all ad spend. Most of this money is bought "
                    "click by click and judged by reports from the platforms that sell it. That is the problem of this part.")
spend = [("$1.14" + NB + "tn", ["global advertising spend in 2025, ", ("+8.8" + PCT, {"bold": True})]),
         ("73" + PCT, ["of it is digital: search, social, video, retail media"]),
         ("€119" + NB + "bn", ["digital advertising in Europe in 2024, ", ("+16" + PCT, {"bold": True})]),
         ("67" + PCT, ["of all ad spend in Europe is digital"])]
for i, (fig, lab) in enumerate(spend):
    y = 1.35 + i * (0.75 + 0.08)
    box(s, 0.5, y, 5.45, 0.75, LIGHT, f"Spend row {i + 1}")
    stat(s, 0.5, y, 5.45, 0.75, fig, [p if isinstance(p, tuple) else (p, {}) for p in lab], f"Spend stat {i + 1}",
         fig_w=2.0, size=12)
card(s, 6.1, 1.35, 3.4, 3.24, "What this means", ["Digital ads are bought click by click",
                                                   "The platforms that sell the ads also report their success",
                                                   "Small measurement errors move billions"], "LuCoins", "Spend card",
     head_h=0.45)
callout(s, 4.72, 0.48, "LuCrosshair", ("The question of this part: ", "who gets the credit for a sale, and is it "
                                                                      "deserved?"), "Spend question", dark=True, size=12)
source(s, "WPP Media, This Year Next Year, mid-year 2025; IAB Europe, AdEx Benchmark 2024")

# M2  Attribution example
s = add_slide("Titel und Inhalt", "Attribution: who gets the credit?",
              notes="Attribution splits the credit for one tracked sale between the ads a customer clicked or saw before "
                    "buying. The journey: a TV spot (not tracked), an Instagram ad click, a YouTube video, then a Google "
                    "search for 'alpenglow chocolate' and an order of EUR 40. Last click gives everything to search, first "
                    "click to Instagram, linear splits equally; data-driven attribution (the GA4 default) estimates shares "
                    "from paths that did and did not end in a purchase. The numbers in the last row are illustrative. TV "
                    "gets nothing in every model, because no click is recorded. And no model asks whether the sale would "
                    "have happened without any ad: that is a causal question for experiments and MMM.")
steps = [("LuTv", "TV spot", "not tracked"), ("LuSmartphone", "Instagram ad", "click"), ("LuVideo", "YouTube video", "view"),
         ("LuSearch", "Google search", "“alpenglow”"), ("LuShoppingCart", "Order", "EUR" + NB + "40")]
bw, aw, ag = 1.5, 0.22, 0.07
for i, (ic, head, sub) in enumerate(steps):
    x = 0.5 + i * (bw + aw + 2 * ag)
    fill = WHITE if i == 0 else (NAVY if i == 4 else LIGHT)
    b = box(s, x, 1.35, bw, 1.05, fill, f"Journey step {i + 1}")
    if i == 0:
        b.line.color.theme_color = NAVY; b.line.dash_style = MSO_LINE.DASH; b.line.width = Pt(1.25)
    ink = WHITE if i == 4 else NAVY
    icon(s, ic, "white" if i == 4 else "navy", x + (bw - 0.4) / 2, 1.42, 0.4, f"journey {i + 1}")
    text(s, x + 0.05, 1.83, bw - 0.1, 0.55, [[(head, {"bold": True, "col": ink, "size": 12})],
                                            [(sub, {"col": ink if i else GREY_TXT, "size": 11})]],
         f"Journey text {i + 1}", align=PP_ALIGN.CENTER, space=0)
    if i < 4:
        arrow(s, x + bw + ag, 1.69, f"Journey arrow {i + 1}")
grid = [("Credit for the order", "TV", "Instagram", "YouTube", "Search"),
        ("Last click", "0" + PCT, "0" + PCT, "0" + PCT, "100" + PCT),
        ("First click", "0" + PCT, "100" + PCT, "0" + PCT, "0" + PCT),
        ("Linear", "0" + PCT, "33" + PCT, "33" + PCT, "33" + PCT),
        ("Data-driven (GA4)", "0" + PCT, "e.g. 25" + PCT, "e.g. 15" + PCT, "e.g. 60" + PCT)]
cws = [2.6, 1.6, 1.6, 1.6, 1.6]
for r_i, row in enumerate(grid):
    y = 2.55 + r_i * 0.42
    x = 0.5
    for c_i, val in enumerate(row):
        w = cws[c_i] - (0.03 if c_i < 4 else 0)
        head_row = r_i == 0
        tv_col = c_i == 1 and not head_row
        box(s, x, y, w, 0.39, ACC if head_row else LIGHT, f"Grid {r_i + 1}-{c_i + 1}")
        col = WHITE if head_row else (ACC if tv_col else NAVY)
        text(s, x + 0.1, y, w - 0.2, 0.39, [[(val, {"bold": head_row or c_i == 0 or tv_col, "col": col,
                                                    "size": 12})]], f"Grid text {r_i + 1}-{c_i + 1}",
             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT if c_i == 0 else PP_ALIGN.CENTER)
        x += cws[c_i]
callout(s, 4.72, 0.48, "LuTriangleAlert", ("All four models share out one sale: ", "none asks if it would have "
                                                                                 "happened without ads."),
        "Attribution catch", dark=True, size=12)
source(s, "Google Analytics 4 documentation; journey and data-driven shares are illustrative")

# M3  Last click
s = add_slide("Titel und Inhalt", "Last click: simple, common and biased",
              notes="Last-click attribution gives all the credit to the last ad clicked before the purchase. It is easy to "
                    "compute and explain, and it is still the default in many shop systems and affiliate contracts. In 2023 "
                    "Google removed first click, linear, time decay and position-based models from Google Ads and GA4, "
                    "leaving data-driven and last click; fewer than 3 % of conversions had used the removed models. Last "
                    "click favours channels that sit at the end of the journey (branded search, retargeting, coupon and "
                    "affiliate sites), which often catch people who had already decided to buy. Upper-funnel media that "
                    "create the demand look worthless. Li and Kannan (2014) show with path data that last click misallocates "
                    "credit and over-credits search.")
cw2 = (9.0 - GAP) / 2
card(s, 0.5, 1.35, cw2, 1.95, "How it works", ["100" + PCT + " of the credit to the last ad clicked before the order",
                                               "Easy to compute, easy to explain",
                                               "GA4 and Google Ads: since 2023 only last click and data-driven are left"],
     "LuMousePointerClick", "LC card", head_h=0.45)
card(s, 0.5 + cw2 + GAP, 1.35, cw2, 1.95, "Who wins, who loses", ["Wins: branded search, retargeting, coupon and "
                                                                  "affiliate sites",
                                                                  "Loses: TV, online video, social, posters",
                                                                  "Ignores: views without a click, offline sales"],
     "LuScale", "WL card", head_h=0.45)
catch = [("The catch: ", "the last click often comes from people who had already decided to buy"),
         ("Example: ", "after a TV spot, a viewer searches “alpenglow”: search gets the sale, TV gets nothing"),
         ("The result: ", "budgets drift to the end of the funnel, and the demand that feeds it slowly dries up")]
for i, para in enumerate(catch):
    y = 3.45 + i * 0.58
    box(s, 0.5, y, 9.0, 0.5, LIGHT, f"Catch row {i + 1}")
    box(s, 0.5, y, 0.07, 0.5, ACC, f"Catch bar {i + 1}")
    text(s, 0.72, y, 8.7, 0.5, [para], f"Catch text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
source(s, "Search Engine Land 2023 (Google Ads and GA4); Li & Kannan 2014 (Journal of Marketing Research)")

# M4  Problems of attribution
s = add_slide("Titel und Inhalt", "Six problems of attribution models",
              notes="1 Correlation, not cause: ads are targeted at people who are likely to buy anyway (retargeting, branded "
                    "search), so credit is not the same as extra sales. 2 Tracking gaps: cookie consent, ad blockers, browser "
                    "limits and Apple's App Tracking Transparency (iOS 14.5, 2021); at launch only about 4 % of US users and "
                    "12 % worldwide allowed tracking (Flurry). 3 Offline is invisible: TV, radio, posters, word of mouth, and "
                    "sales in shops. 4 Every platform counts its own conversions with its own rules and windows, so the sum "
                    "of the reports is larger than the real number of orders. 5 Fraud and poaching: affiliates and ad "
                    "networks can 'claim' sales that would have happened anyway; Uber found fake clicks attached to organic "
                    "installs. 6 Many devices: one person on phone, tablet and laptop looks like three customers.")
tw3 = (9.0 - GAP * 2) / 3
probs = [("LuShuffle", "Correlation only", ["Ads target people who buy anyway", "Credit is not extra sales"]),
         ("LuEyeOff", "Tracking gaps", ["Cookie consent, ad blockers, Apple's ATT",
                                        "2021: 4" + PCT + " of US iPhone users allowed tracking"]),
         ("LuTv", "Offline is invisible", ["TV, radio, posters, word of mouth", "Sales in shops have no click"]),
         ("LuCopy", "Double counting", ["Each platform counts its own orders", "Sum of reports > real orders"]),
         ("LuBug", "Click fraud", ["Affiliates claim sales that come anyway", "Uber: fake clicks on free installs"]),
         ("LuSmartphone", "Many devices", ["Phone on the sofa, laptop at work", "One buyer looks like three"])]
for i, (ic, head, lines) in enumerate(probs):
    x = 0.5 + (i % 3) * (tw3 + GAP)
    y = 1.35 + (i // 3) * (1.62 + 0.1)
    tile(s, x, y, tw3, 1.62, ic, head, [[(ln, {})] for ln in lines], f"Problem {i + 1}", size=12)
callout(s, 4.8, 0.42, "LuFlaskConical", ("The remedy: ", "experiments (ads off in some regions) and MMM, which needs "
                                                         "no tracking."), "Remedy", size=12)
source(s, "Gordon et al. 2019 (Marketing Science); Flurry Analytics 2021; WARC 2019 (Uber)")

# M5  Real-world tests
s = add_slide("Titel und Inhalt", "When the ads were switched off",
              notes="eBay (Blake, Nosko and Tadelis 2015, Econometrica): when eBay stopped paying for brand keywords, 99.5 % "
                    "of that traffic came back through the free search results; in a test across US regions, non-brand "
                    "search mainly reached people who would have bought anyway, and the return was far below what the "
                    "click reports showed. Uber: Kevin Frisch turned off $100 million of $150 million yearly performance "
                    "ads and saw no change in rider app installs; much of the spend was fraud. P&G cut $200 million of "
                    "digital ads in 2017 that did not reach real people, moved money to TV, audio and e-commerce, and its "
                    "reach rose about 10 %. Airbnb cut marketing from $1.62 billion (2019) to $0.55 billion (2020), mostly "
                    "performance ads, and kept about 95 % of its traffic; Covid is a confound, but Airbnb kept the shift. "
                    "Gordon et al. (2019) compared 15 Facebook experiments with observational methods, which often missed "
                    "the true lift, mostly overstating it.")
tests = [("99.5" + PCT, [("eBay: ", {"bold": True}), ("brand-keyword traffic that came back through the free links "
                                                       "when the ads were paused", {})]),
         ("$100" + NB + "m", [("Uber: ", {"bold": True}), ("of $150" + NB + "m yearly performance ads switched off, with no "
                                                            "change in app installs", {})]),
         ("$200" + NB + "m", [("P&G: ", {"bold": True}), ("of digital ads cut in 2017; reach rose about 10" + PCT, {})]),
         ("95" + PCT, [("Airbnb: ", {"bold": True}), ("of web traffic kept in 2020 after cutting marketing by more than "
                                                      "half", {})]),
         ("15", [("Facebook experiments: ", {"bold": True}), ("click-based methods often missed the true lift, mostly "
                                                              "overstating it", {})])]
for i, (fig, lab) in enumerate(tests):
    y = 1.35 + i * (0.62 + 0.07)
    box(s, 0.5, y, 9.0, 0.62, LIGHT, f"Test row {i + 1}")
    stat(s, 0.5, y, 9.0, 0.62, fig, lab, f"Test stat {i + 1}", fig_w=2.25, size=12)
callout(s, 4.85, 0.4, "LuLightbulb", ("Lesson: ", "test what the ads add; do not trust the click report."), "Test lesson",
        dark=True, size=12)
source(s, "Blake, Nosko & Tadelis 2015; WARC 2019; Adweek 2018; Campaign Asia 2021; Gordon et al. 2019", y=5.29)

# M6  In-class discussion: attribution
s = add_slide("Titel und Inhalt", "Discuss: who sold the chocolate?",
              notes="15 minutes: 10 in pairs, 5 in plenary. Suggested answers. Q1: every platform counts its own clicks and "
                    "views with its own window, so one order is claimed two or three times; some orders would have come "
                    "anyway. Q2: no. Branded search catches people TV made curious (the mediation in case task B5); a last-"
                    "click ROAS of 12 says nothing about extra sales; cutting TV may later reduce searches. Q3: switch "
                    "branded search off in some German regions or weeks and compare orders with the others (geo test, like "
                    "eBay); or a holdout group. Q4: MMM sees TV, prices, holidays and weather without tracking, but it "
                    "needs years of data, sees only weekly totals, and cannot split credit for one customer.")
stepper(s, 0.5, 1.35, 9.0, 0.5, ["Pairs: 10 minutes", "Plenary: 5 minutes"], "Discussion steps", size=12)
box(s, 0.5, 2.0, 3.3, 3.12, LIGHT, "Reports panel")
text(s, 0.65, 2.07, 3.0, 0.42, [[("Alpenglow Germany, November", {"bold": True, "col": NAVY, "size": 12})]],
     "Reports heading", anchor=MSO_ANCHOR.MIDDLE)
reports = [("Google Ads (last click)", "3,900", False), ("Meta (click or view)", "2,600", False),
           ("Affiliate sites", "1,300", False), ("Sum of the reports", "7,800", True),
           ("Orders in the web shop", "6,000", True)]
for i, (lab, val, strong) in enumerate(reports):
    y = 2.55 + i * 0.49 + (0.1 if i >= 3 else 0)
    if i == 3:
        box(s, 0.65, y - 0.06, 3.0, 0.02, NAVY, "Reports rule")
    col = ACC if i == 4 else NAVY
    text(s, 0.65, y, 2.2, 0.42, [[(lab, {"bold": strong, "col": col, "size": 12})]], f"Report label {i + 1}",
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.8, y, 0.85, 0.42, [[(val, {"bold": True, "col": col, "size": 12})]], f"Report value {i + 1}",
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
numbered(s, 3.95, 2.0, 5.55, [
    ("Why ", "do the reports add up to 7,800 orders when the shop had 6,000?"),
    ("Branded search ", "shows a last-click ROAS of 12, TV of 2. The CFO wants to move TV money to search. Your answer?"),
    ("Design a test: ", "how would you find out what branded search really adds?"),
    ("MMM: ", "what can it see that attribution cannot, and what can it not see?")],
    "Attribution questions", row_h=0.72, gap=0.08, size=12)
source(s, "Illustrative numbers; real-world cases on the previous slide", y=5.27)

# ---------------------------------------------------------------- Part 6 -----
divider("Reading: Analytics for Marketers", "Part 6 of 7" + NB + "·" + NB + "Fantini & Narayandas (HBR 2023) and an "
        "in-class discussion", "Reading for this session: Fantini, F. & Narayandas, D. (2023). Analytics for Marketers: "
        "when to rely on algorithms and when to trust your gut. Harvard Business Review, May–June 2023 (on Canvas).")

# R1  Three approaches
s = add_slide("Titel und Inhalt", "Three approaches to analytics",
              notes="The article sorts analytics by who decides. Descriptive: humans read aggregated past data (business "
                    "intelligence, dashboards) and decide; cheap, but coarse and dependent on the manager's intuition and "
                    "biases. Predictive: models estimate what will happen for different inputs (regression, elasticities, "
                    "forecasts, segments); humans choose. Prescriptive: machines decide within objectives, rules and "
                    "constraints set by managers, learning from many low-cost experiments and modelling uncertainty and "
                    "costs; expensive to set up. MMM is mostly predictive; a budget optimiser on top makes it prescriptive.")
stepper(s, 0.5, 1.35, 9.0, 0.6, ["Descriptive:\vhumans decide", "Predictive: machines\vpredict, humans choose",
                                 "Prescriptive:\vmachines decide"], "Approach steps", size=12)
tw3 = (9.0 - GAP * 2) / 3
approaches = [("LuChartColumn", "What happened?", ["Dashboards of aggregated past data", "Managers add experience and gut feeling",
                                                    "Low pain, low gain"]),
              ("LuTrendingUp", "What if …?", ["Regression, elasticities, forecasts", "Finer, tactical decisions",
                                                             "Few variables, uncertain inputs"]),
              ("LuBot", "What to do?", ["Optimises towards goals and rules set by managers", "Learns from many cheap experiments",
                                               "Costly to build and run"])]
for i, (ic, head, items) in enumerate(approaches):
    card(s, 0.5 + i * (tw3 + GAP), 2.15, tw3, 2.45, head, items, ic, f"Approach card {i + 1}", head_h=0.45)
callout(s, 4.75, 0.5, "LuSlidersHorizontal", ("MMM today: ", "predictive. With a budget optimiser on top it becomes "
                                                             "prescriptive."), "MMM approach", size=12)
source(s, "Fantini & Narayandas 2023 (Harvard Business Review)")

# R2  When to use which
s = add_slide("Titel und Inhalt", "When to use which approach",
              notes="Two factors decide: how relevant and plentiful the data are, and the business case (how much profit "
                    "better decisions can unlock). Little data and high uncertainty: descriptive gives direction. Many "
                    "repeated, granular decisions with rich data: prescriptive pays for its cost. In between, and for long "
                    "horizons, noisy granular data or small gains from extreme optimisation: predictive. Strategy and "
                    "innovation stay with humans, because asking the right question matters more than an accurate answer.")
box(s, 0.5, 1.35, 4.6, 3.75, LIGHT, "Map area")
box(s, 0.95, 1.55, 0.03, 3.0, NAVY, "Map y axis")
box(s, 0.95, 4.52, 4.0, 0.03, NAVY, "Map x axis")
text(s, 0.55, 1.5, 0.35, 3.0, [[("business case", {"size": 11, "col": NAVY, "bold": True})]], "Map y label",
     anchor=MSO_ANCHOR.MIDDLE)
s.shapes[-1].rotation = 270
s.shapes[-1].left, s.shapes[-1].top, s.shapes[-1].width, s.shapes[-1].height = I(-0.85), I(2.88), I(3.0), I(0.35)
text(s, 0.95, 4.6, 4.0, 0.35, [[("relevant data, decisions repeated often", {"size": 11, "col": NAVY, "bold": True})]],
     "Map x label", align=PP_ALIGN.CENTER)
zones = [("Descriptive", "little data, high uncertainty", 1.15, 3.6, LIGHT, NAVY, "C8E8F4"),
         ("Predictive", "the middle ground", 2.05, 2.75, ACC, WHITE, None),
         ("Prescriptive", "rich data, many decisions", 2.95, 1.75, NAVY, WHITE, None)]
for j, (name, sub, x, y, fill, ink, custom) in enumerate(zones):
    b = box(s, x, y, 1.95, 0.75, custom or fill, f"Zone {j + 1}", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.1, y, 1.75, 0.75, [[(name, {"bold": True, "col": ink, "size": 12})], [(sub, {"col": ink, "size": 11})]],
         f"Zone text {j + 1}", anchor=MSO_ANCHOR.MIDDLE, space=0)
rows = [("LuCompass", ("Humans lead: ", "strategy, innovation, entering a new market")),
        ("LuChartLine", ("Predictive fits: ", "planning, noisy customer data, small gains from fine-tuning")),
        ("LuSettings", ("Prescriptive pays: ", "prices, stock, marketing spend, decided often with rich data"))]
for i, (ic, para) in enumerate(rows):
    y = 1.35 + i * (1.2 + 0.075)
    box(s, 5.25, y, 4.25, 1.2, LIGHT, f"Use row {i + 1}")
    icon(s, ic, "navy", 5.45, y + 0.38, 0.44, f"use {i + 1}")
    text(s, 6.05, y, 3.35, 1.2, [para], f"Use text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
source(s, "Fantini & Narayandas 2023 (Harvard Business Review); simplified illustration")

# R3  Event Network case
s = add_slide("Titel und Inhalt", "Case: markdowns at Event Network",
              notes="Event Network runs gift shops in museums, zoos and aquariums in the US and Canada, with more than "
                    "100,000 SKUs. Descriptive: discount the SKUs with the highest coverage ratio (days of stock left) until "
                    "the markdown budget is spent; ignores customers. Predictive: regress volume on price by category, store "
                    "and week to get elasticities. With an elasticity of −2, a 10 % price cut raises units by 20 %: 100 units "
                    "at $10 ($1,000) become 120 units at $9 ($1,080), 8 % more revenue. Price explained only 10 to 20 % of "
                    "sales, yet the crude model beat the descriptive rule. Prescriptive: machine learning and automated "
                    "optimisation over many data sources decide which products to discount, when and by how much.")
tw3 = (9.0 - GAP * 2) / 3
ev = [("LuChartColumn", "1 Descriptive", ["Discount SKUs with the most days of stock", "Until the markdown budget is spent",
                                          "Ignores customers and context"]),
      ("LuTrendingUp", "2 Predictive", ["Regression of volume on price", "Elasticity by category, store and week",
                                        "Crude, but better than the rule"]),
      ("LuBot", "3 Prescriptive", ["Machine learning plus optimisation", "Many data sources, 100,000+ SKUs",
                                   "Learns and improves over time"])]
for i, (ic, head, items) in enumerate(ev):
    card(s, 0.5 + i * (tw3 + GAP), 1.35, tw3, 2.15, head, items, ic, f"EN card {i + 1}", head_h=0.45)
tw4 = (9.0 - GAP * 3) / 4
nums = [("−2", "price elasticity"), ("+20" + PCT, "units after a 10" + PCT + " price cut"),
        ("+8" + PCT, "revenue: $1,000 → $1,080"), ("10–20" + PCT, "of sales explained by price")]
for i, (fig, lab) in enumerate(nums):
    x = 0.5 + i * (tw4 + GAP)
    box(s, x, 3.65, tw4, 1.05, LIGHT, f"EN stat {i + 1}")
    box(s, x, 3.65, 0.07, 1.05, ACC, f"EN stat bar {i + 1}")
    text(s, x + 0.2, 3.68, tw4 - 0.3, 0.5, [[(fig, {"size": FIG, "bold": True, "col": NAVY, "head": True})]],
         f"EN stat figure {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 4.18, tw4 - 0.3, 0.48, [lab], f"EN stat label {i + 1}", size=11.5, space=0)
callout(s, 4.8, 0.42, "LuLightbulb", ("Lesson: ", "a crude model can still beat a rule of thumb."), "EN lesson", size=12)
source(s, "Fantini & Narayandas 2023 (Harvard Business Review)", y=5.27)

# R4  Accuracy versus business value
s = add_slide("Titel und Inhalt", "Accuracy is not the goal",
              notes="Data scientists tend to maximise accuracy; business scientists maximise impact by weighing the cost "
                    "of each kind of error. A false positive (predicted win, no win) wastes sales and marketing effort; a "
                    "false negative (predicted loss, but it would have been a win) loses the opportunity. A model that cuts "
                    "false positives but misses many real chances can be accurate and still unprofitable. Session 3 and "
                    "task C7 of the case study do exactly this: choose the churn threshold by profit. The closing quote is "
                    "the article's main message about the manager's role.")
cw2 = (9.0 - GAP) / 2
card(s, 0.5, 1.35, cw2, 1.75, "Data scientists", ["Aim: the most accurate model", "Judge by statistical fit"],
     "LuSigma", "DS card", head_h=0.45)
card(s, 0.5 + cw2 + GAP, 1.35, cw2, 1.75, "Business scientists", ["Aim: the most profitable decision",
                                                                  "Judge by the cost of each error"],
     "LuCoins", "BS card", head_h=0.45)
errs = [("False positive: ", "predicted a win that does not come: wasted sales and marketing effort"),
        ("False negative: ", "predicted a loss that would have been a win: a lost opportunity"),
        ("In this course: ", "Session 3 and case task C7 set the churn threshold by profit, not accuracy")]
for i, para in enumerate(errs):
    y = 3.25 + i * 0.52
    box(s, 0.5, y, 9.0, 0.46, LIGHT, f"Error row {i + 1}")
    box(s, 0.5, y, 0.07, 0.46, ACC, f"Error bar {i + 1}")
    text(s, 0.72, y, 8.7, 0.46, [para], f"Error text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
callout(s, 4.78, 0.47, "LuMessagesSquare", ("The manager's new role: ", "“from the person who has all the answers to the "
                                                                       "one who asks the right questions”"),
        "Quote", dark=True, size=11.5)
source(s, "Fantini & Narayandas 2023 (Harvard Business Review)")

# D1  In-class assignment: the task
s = add_slide("Titel und Tabelle", "In-class discussion: six Alpenglow decisions",
              notes="25 minutes: 15 in groups of 3 to 4, 10 in plenary. Suggested answers. 1 Media budget: predictive (MMM "
                    "estimates response curves; humans decide once a year, stakes are high, data moderate); an optimiser "
                    "can propose, managers approve. 2 Search bids: prescriptive (thousands of small, repeated decisions "
                    "with rich data; automated bidding). 3 Polish Christmas price: predictive (elasticity from regression; "
                    "the brand manager decides with positioning in mind). 4 Retention offer: predictive moving to "
                    "prescriptive (score plus profit threshold, can run automatically every month). 5 Swiss retail entry: "
                    "descriptive and human judgement (one-off, little data, high uncertainty, strategy). 6 Easter "
                    "markdowns: prescriptive (the Event Network situation: many SKUs and stores, repeated, rich data).")
rows = [("", "Decision", "How often", "Data available"),
        ("1", "Split next year's media budget over 6 countries and 5 channels", "once a year", "3 years of weekly sales and spend"),
        ("2", "Set the bids for paid-search keywords", "every day", "clicks, costs and orders per keyword"),
        ("3", "Price of the Christmas gift box in Poland", "once a year", "156 weeks of prices and sales"),
        ("4", "Choose which Club members get a retention offer", "every month", "6,000 subscribers and their history"),
        ("5", "Enter Swiss supermarkets, beyond the web shop", "once", "almost none"),
        ("6", "Discount unsold Easter chocolate after Easter", "every year, 200 SKUs × stores", "sales and stock per SKU and store")]
tbl = table(s, rows, [0.45, 4.15, 1.75, 2.65], "Decision table", size=12, top=1.35)
for r_i in range(1, len(rows)):
    tbl.cell(r_i, 0).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
text(s, 0.5, 4.95, 9.0, 0.5, [[("Your task: ", {"bold": True, "col": NAVY}),
                               ("for each decision, choose descriptive, predictive or prescriptive, and say who decides.",
                                {})]], "Task line", anchor=MSO_ANCHOR.MIDDLE, size=12)

# D2  In-class assignment: questions and timing
s = add_slide("Titel und Inhalt", "Discuss in your group",
              notes="Expected points. Q2: today the MMM is predictive; a budget optimiser makes it prescriptive, but the "
                    "yearly budget is high-stakes, rare and politically loaded, so managers should approve. Q3: a crude "
                    "model is good enough when it beats the current rule and its errors are cheap; Event Network's case. "
                    "Q4: depends on costs: with a EUR 5 offer and EUR 80 margin, a missed churner (false negative) costs "
                    "more than an unnecessary offer. Q5: strategic questions, for example which brand Alpenglow wants to "
                    "be in Poland, or whether to enter retail at all.")
stepper(s, 0.5, 1.35, 9.0, 0.55, ["Groups of 3–4:\v15 minutes", "Plenary:\v10 minutes", "Hand in: one\vphoto of your table"],
        "Timing steps", size=12)
numbered(s, 0.5, 2.05, 9.0, [
    ("Classify ", "the six decisions with the article's two criteria: relevant data and the business case"),
    ("Alpenglow's MMM: ", "predictive or prescriptive? Should a machine decide the yearly budget?"),
    ("Crude models: ", "Event Network's price model explained 10–20" + PCT + " of sales. When is that good enough?"),
    ("Errors: ", "for the retention offer, which costs more: a false positive or a false negative?"),
    ("The manager's job: ", "write one question for Alpenglow's CMO that no model can answer")],
    "Questions", row_h=0.56, gap=0.08, size=12)
source(s, "Reading: Fantini & Narayandas 2023, Analytics for Marketers, Harvard Business Review (on Canvas)", y=5.3)

# ---------------------------------------------------------------- Part 7 -----
divider("MMM in this course", "Part 7 of 7" + NB + "·" + NB + "how the sessions build the skills, the group case study",
        "Every session adds one piece of the model. By Session 5 you can build and judge a full MMM.")

# 20  Sessions
s = add_slide("Titel und Inhalt", "Five sessions, one model",
              notes="Session 1: ordinary least squares regression and how to read model fit, the base of every MMM. "
                    "Session 2: the MMM itself, with log-log elasticities, dummies for holidays, non-linear effects, "
                    "moderation (does an effect depend on another variable?) and mediation (does TV work through search?), "
                    "adstock and saturation, and diagnostics. Session 3: logistic regression for yes/no outcomes such as "
                    "purchase or churn. Session 4: ARIMA and ARIMAX for the sales baseline and forecasts. Session 5: "
                    "Bayesian MMM tools and budget optimisation.")
sess = [("Introduction and linear regression", "OLS, model fit and interpretation", "the base regression"),
        ("Advanced regression", "log-log elasticities, dummies, non-linear effects, moderation and mediation, adstock "
                                "and saturation, diagnostics", "the MMM itself"),
        ("Logistic regression", "binary outcomes: purchase, churn; odds ratios", "customer-level questions"),
        ("ARIMA", "trend, seasonality, ARIMA and ARIMAX forecasts", "the sales baseline"),
        ("Advanced MMM tools", "Bayesian MMM, budget optimisation", "the full workflow")]
rh, rg = 0.74, 0.085
for i, (topic, methods, role) in enumerate(sess):
    y = 1.35 + i * (rh + rg)
    box(s, 0.5, y, 9.0, rh, LIGHT, f"Session row {i + 1}")
    box(s, 0.5, y, 0.07, rh, ACC, f"Session bar {i + 1}")
    badge(s, 0.72, y + (rh - 0.42) / 2, i + 1, f"Session number {i + 1}", d=0.42, size=14)
    text(s, 1.3, y, 2.25, rh, [[(topic, {"bold": True, "col": NAVY})]], f"Session topic {i + 1}", anchor=MSO_ANCHOR.MIDDLE,
         size=12)
    text(s, 3.6, y, 3.6, rh, [methods], f"Session methods {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=12)
    text(s, 7.3, y, 2.1, rh, [[("→ " + role, {"bold": True, "col": ACC})]], f"Session role {i + 1}",
         anchor=MSO_ANCHOR.MIDDLE, size=12)

# 21  Group case study
s = add_slide("Titel und Inhalt", "The group case study",
              notes="The case study applies Sessions 1 to 4 to new data: Alpenglow's own web shop in Austria, Germany and "
                    "Switzerland, plus its subscription club. Each block of the course answers one of the Head of "
                    "E-commerce's questions. Details, data and assessment criteria are in the brief in "
                    "assignments/group-project/ and on Canvas.")
stepper(s, 0.5, 1.35, 9.0, 0.66, ["S1: what drives\vweekly orders?", "S2: price and\vadvertising", "S3: who will\vcancel?",
                                  "S4: Q1 2026\vbaseline"], "Case steps", size=12)
tw3 = (9.0 - GAP * 2) / 3
case = [("LuCandy", "The brief", ["Alpenglow's web shop in AT, DE and CH", "Four questions from the Head of E-commerce"]),
        ("LuDatabase", "The data", ["Weekly shop data, 3 countries × 156 weeks", "Club subscribers", "The Q1 2026 plan"]),
        ("LuFileText", "You deliver", ["A Quarto report with all code", "Groups of at most 4", "30" + PCT + " of the grade"])]
for i, (ic, head, items) in enumerate(case):
    card(s, 0.5 + i * (tw3 + GAP), 2.2, tw3, 2.0, head, items, ic, f"Case card {i + 1}", head_h=0.45)
callout(s, 4.38, 0.75, "LuBriefcase", ("Every number from code: ", "the report renders from a fresh copy of your "
                                                                  "project folder."), "Case rule", dark=True, size=12)
text(s, 0.5, 5.2, 9.0, 0.32, [[("Brief, data and assessment criteria: ", {"bold": True, "col": NAVY, "size": 12}),
                               ("assignments/group-project/ and Canvas", {"font": MONO, "col": NAVY, "size": 12})]],
     "Case where", anchor=MSO_ANCHOR.MIDDLE)

# 22  Key takeaways
s = add_slide("Titel und Inhalt", "Key takeaways",
              notes="If students remember one thing: decisions are made at the margin, and an MMM is only as good as the "
                    "variation in its data and the experiments that check it.")
numbered(s, 0.5, 1.35, 9.0, [
    ("MMM ", "explains weekly sales with marketing and other drivers and splits them into base and incremental"),
    ("Adstock and saturation ", "capture carryover and diminishing returns; decide on marginal, not average, ROAS"),
    ("Benchmarks: ", "advertising elasticities around 0.1 short term, larger price effects, differences by country"),
    ("Budgets follow demand: ", "watch endogeneity; validate on a holdout and calibrate with experiments"),
    ("Open-source Bayesian tools ", "make it doable in Python: you will build one")], "Takeaways", row_h=0.64,
    gap=0.09, size=12)
callout(s, 5.0, 0.55, "LuWallet", ("The question stays the same: ", "where should the next euro go?"), "Final question",
        dark=True, size=12)

# 23  Further reading
s = add_slide("Titel und Inhalt", "Further reading",
              notes="Start with Chan and Perry (short, free) and Jin et al. (free) for the modern model, Sethuraman et al. "
                    "for benchmarks, and Gordon et al. for why experiments matter. Hanssens et al. is the reference book "
                    "(WU library). The full annotated list is in instructor/resources/mmm-resource-guide.md.")
refs_l = [("Hanssens, D. M., Parsons, L. J. & Schultz, R. L. (2001). ", "Market Response Models. 2nd ed. Kluwer.", None),
          ("Sethuraman, R., Tellis, G. J. & Briesch, R. A. (2011). ", "How well does advertising work? Journal of "
           "Marketing Research 48(3), 457–471.", "https://doi.org/10.1509/jmkr.48.3.457"),
          ("Gordon, B. R., Zettelmeyer, F., Bhargava, N. & Chapsky, D. (2019). ", "A comparison of approaches to "
           "advertising measurement. Marketing Science 38(2), 193–225.", "https://doi.org/10.1287/mksc.2018.1135"),
          ("Chan, D. & Perry, M. (2017). ", "Challenges and opportunities in media mix modeling. Google.",
           "https://research.google/pubs/pub45998/")]
refs_r = [("Jin, Y., Wang, Y., Sun, Y., Chan, D. & Koehler, J. (2017). ", "Bayesian methods for media mix modeling "
           "with carryover and shape effects. Google.", "https://research.google/pubs/pub46001/"),
          ("Runge, J., Skokan, I., Zhou, G. & Pauwels, K. (2024). ", "Packaging up media mix modeling: Robyn's "
           "open-source approach. arXiv 2403.14674.", "https://arxiv.org/abs/2403.14674"),
          ("Google Meridian. ", "Documentation and code (Python).", "https://github.com/google/meridian"),
          ("PyMC Labs. PyMC-Marketing. ", "Documentation and code (Python).",
           "https://github.com/pymc-labs/pymc-marketing")]
cw2 = (9.0 - GAP) / 2
for k, (head, ic, refs) in enumerate([("Foundations and evidence", "LuBookOpen", refs_l),
                                      ("Modern MMM and tools", "LuCode", refs_r)]):
    x = 0.5 + k * (cw2 + GAP)
    box(s, x, 1.35, cw2, 0.45, ACC, f"Reading {k + 1} header")
    text(s, x + 0.15, 1.35, cw2 - 0.75, 0.45, [head], f"Reading {k + 1} heading", size=HEAD, col=WHITE, bold=True,
         head=True, anchor=MSO_ANCHOR.MIDDLE)
    icon(s, ic, "white", x + cw2 - 0.48, 1.4, 0.34, f"reading {k + 1}")
    box(s, x, 1.8, cw2, 3.75, LIGHT, f"Reading {k + 1} body")
    paras = []
    for lead, rest, url in refs:
        paras.append([(lead, {"bold": True, "col": NAVY}), (rest, {})])
        if url:
            paras.append([(url.replace("https://", ""), {"link": url, "col": ACC})])
    text(s, x + 0.18, 1.92, cw2 - 0.33, 3.55, paras, f"Reading {k + 1} text", size=11, space=5)

# 24  Closing contact card
add_slide("Abschlussfolie Kontakt", notes="Questions? Contact details on the card; office hours via the booking link.")

deck.save(OUT)
