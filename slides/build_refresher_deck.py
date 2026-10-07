"""Build the deck "Statistics refresher" (Session 1, before regression) on the WU template (polish-slides skill).

Five parts on one running example (a chocolate kiosk, five weeks, small enough for a calculator): scale types;
mean, median, variance, standard deviation and standard error; standardisation, normal and standard normal
distribution; covariance and correlation with the bridge to regression; and the same calculations by hand in
Python (polars, plotnine) with a check against the built-in functions, plus an in-class exercise.
Every number comes from make_refresher_figures.py (slides/figures/ref_numbers.txt), which also draws the charts.

usage: python slides/build_refresher_deck.py TEMPLATE.pptx OUT.pptx
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from wu_deck import (A, ACC, BODY, FIG, GAP, HEAD, INK, LIGHT, MONO, NAVY, NB, WHITE, Deck, I, box, card,
                     code_block, color, icon, ph, set_runs, stepper, text)

TEMPLATE, OUT = sys.argv[1:3]
FIGS = Path(__file__).resolve().parent / "figures"
FOOTER = "International Marketing Analytics" + NB + "·" + NB + "WT" + NB + "2026/27"
deck = Deck(TEMPLATE, FOOTER)
add_slide, L = deck.add_slide, deck.L
GREY_TXT = "595959"
PCT = NB + "%"


# ---------------------------------------------------------------- helpers ----
def divider(title, part, notes):
    s = add_slide("Kapitelfolie", title, notes=notes)
    set_runs(ph(s, 11).text_frame.paragraphs[0], part)
    return s


def badge(slide, x, y, num, name, d=0.34, fill=NAVY, size=12):
    b = box(slide, x, y, d, d, fill, name, shape=MSO_SHAPE.OVAL)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(num); r.font.size = Pt(size); r.font.bold = True
    color(r.font.color, WHITE)


def numbered(slide, x, y, w, items, name, row_h=0.62, gap=0.1, size=BODY):
    """Numbered rows: bg2 band, navy number badge, text with a bold lead-in."""
    for i, para in enumerate(items):
        ry = y + i * (row_h + gap)
        box(slide, x, ry, w, row_h, LIGHT, f"{name} row {i + 1}")
        badge(slide, x + 0.15, ry + (row_h - 0.34) / 2, i + 1, f"{name} number {i + 1}")
        text(slide, x + 0.65, ry, w - 0.8, row_h, [para], f"{name} text {i + 1}", size=size, anchor=MSO_ANCHOR.MIDDLE)


def callout(slide, y, h, icon_name, para, name, dark=False, x=0.5, w=9.0, size=BODY):
    box(slide, x, y, w, h, NAVY if dark else LIGHT, f"{name} box")
    icon(slide, icon_name, "white" if dark else "navy", x + 0.22, y + (h - 0.4) / 2, 0.4, name)
    if dark:
        lead, rest = para
        para = [(lead, {"bold": True, "col": WHITE}), (rest, {"col": WHITE})]
    text(slide, x + 0.82, y, w - 1.0, h, [para], f"{name} text", anchor=MSO_ANCHOR.MIDDLE, size=size)


def tile(slide, x, y, w, h, icon_name, head, lines, name, size=BODY):
    """Icon at the top left, bold navy heading, plain lines underneath."""
    box(slide, x, y, w, h, LIGHT, f"{name} tile")
    icon(slide, icon_name, "navy", x + 0.15, y + 0.15, 0.4, name)
    text(slide, x + 0.65, y + 0.15, w - 0.75, 0.4, [[(head, {"bold": True, "col": NAVY, "size": HEAD, "head": True})]],
         f"{name} heading", anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.15, y + 0.68, w - 0.3, h - 0.75, lines, f"{name} text", space=4, size=size)


def stat(slide, x, y, w, h, fig, label, name, fig_w=1.25, size=BODY):
    """Stat row: thin accent bar, 28 pt navy figure, explanation to its right."""
    box(slide, x, y, 0.07, h, ACC, f"{name} bar")
    text(slide, x + 0.2, y, fig_w, h, [[(fig, {"size": FIG, "bold": True, "col": NAVY, "head": True})]],
         f"{name} figure", anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.2 + fig_w, y, w - fig_w - 0.25, h, [label], f"{name} label", anchor=MSO_ANCHOR.MIDDLE, size=size)


def source(slide, txt, y=5.27, x=0.5, w=9.0):
    text(slide, x, y, w, 0.28, [[("Sources: " + txt, {"col": GREY_TXT, "size": 11})]], "Sources",
         anchor=MSO_ANCHOR.MIDDLE)


def picture(slide, fname, x, y, w, h, alt):
    pic = slide.shapes.add_picture(str(FIGS / fname), I(x), I(y), I(w), I(h))
    pic.name = "Chart " + fname
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def arrow(slide, x, y, name, w=0.22, h=0.36):
    box(slide, x, y, w, h, ACC, name, shape=MSO_SHAPE.CHEVRON)


def table(slide, rows, widths, name, size=BODY, head_size=HEAD, top=1.35, bold_first_col=True, row_margin=45000,
          mono_rows=()):
    gf = ph(slide, 1).insert_table(len(rows), len(rows[0]))
    gf.name = name
    tbl = gf.table
    for i, w in enumerate(widths):
        tbl.columns[i].width = I(w)
    gf.left, gf.top, gf.width = I(0.5), I(top), I(sum(widths))
    for r_i, row in enumerate(rows):
        for c_i, txt in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            p = cell.text_frame.paragraphs[0]
            set_runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(head_size if r_i == 0 else size)
                if c_i == 0 and r_i > 0 and bold_first_col:
                    r.font.bold = True
            cell.margin_top = cell.margin_bottom = Emu(row_margin)
    return tbl


def link(t, url):
    return (t, {"link": url, "col": ACC})




def rows_with_bar(slide, x, y, w, items, name, row_h=0.62, gap=0.08, size=12):
    for i, para in enumerate(items):
        ry = y + i * (row_h + gap)
        box(slide, x, ry, w, row_h, LIGHT, f"{name} row {i + 1}")
        box(slide, x, ry, 0.07, row_h, ACC, f"{name} bar {i + 1}")
        text(slide, x + 0.22, ry, w - 0.32, row_h, [para], f"{name} text {i + 1}", anchor=MSO_ANCHOR.MIDDLE, size=size)


def grid(slide, rows, widths, x, y, name, row_h=0.36, size=11.5, head_fill=ACC, first_bold=True, last_bold=False,
         align=None):
    """Drawn table: header in accent1, body rows in bg2; optional bold last row (sums)."""
    for r_i, row in enumerate(rows):
        cx = x
        last = last_bold and r_i == len(rows) - 1
        for c_i, val in enumerate(row):
            w = widths[c_i]
            head_row = r_i == 0
            box(slide, cx, y + r_i * (row_h + 0.03), w - 0.03, row_h, head_fill if head_row else LIGHT,
                f"{name} {r_i + 1}-{c_i + 1}")
            col = WHITE if head_row else (NAVY if (c_i == 0 and first_bold) or last else INK)
            al = align[c_i] if align else PP_ALIGN.LEFT
            text(slide, cx + 0.1, y + r_i * (row_h + 0.03), w - 0.23, row_h,
                 [[(val, {"bold": head_row or last or (c_i == 0 and first_bold), "col": col, "size": size})]],
                 f"{name} text {r_i + 1}-{c_i + 1}", anchor=MSO_ANCHOR.MIDDLE, align=al)
            cx += w


R = PP_ALIGN.RIGHT
L_ = PP_ALIGN.LEFT

# ================================================================ slides ======
# 1  Title
kontakt = next(sh for sh in L["Titelfolie Kontakt"].shapes if sh.name == "Kontaktdaten")
for para in kontakt.text_frame.paragraphs:
    if "Institute for International Business" in para.text:
        blank = copy.deepcopy(para._p)
        for r in blank.findall(f"{{{A}}}r"):
            blank.remove(r)
        end = etree.SubElement(blank, f"{{{A}}}endParaRPr"); end.set("lang", "en-GB"); end.set("sz", "1200")
        para._p.addnext(blank)
        break
s = add_slide("Titelfolie Kontakt", footer=None,
              notes="A refresher before regression: the eight ideas every later session uses. One small running example "
                    "(a chocolate kiosk) so that everything can be followed with a calculator, then the same calculations "
                    "by hand in Python, in Positron.")
set_runs(ph(s, 0).text_frame.paragraphs[0], "International Marketing Analytics")
set_runs(ph(s, 1).text_frame.paragraphs[0], "Statistics refresher")
set_runs(ph(s, 12).text_frame.paragraphs[0], "Session 1" + NB + "·" + NB + "WT" + NB + "2026/27")

# 2  Today
s = add_slide("Titel und Inhalt", "Today",
              notes="Four blocks of theory, each on the kiosk numbers, then Python: we calculate mean, standard deviation, "
                    "z-scores, covariance and correlation by hand in code and check them against the built-in functions.")
stepper(s, 0.5, 1.35, 9.0, 0.6, ["Scale types", "Centre and\vspread", "Normal\vdistribution", "Covariance and\vcorrelation",
                                 "Python in\vPositron"], "Plan steps", size=12)
cw2 = (9.0 - GAP) / 2
card(s, 0.5, 2.1, cw2, 2.55, "The concepts", [
    "Scale types: what you may calculate",
    "Mean, median, variance, standard deviation",
    "Standard error and standardisation (z)",
    "Normal and standard normal distribution",
    "Covariance and correlation"], "LuSigma", "Concepts card", head_h=0.45)
card(s, 0.5 + cw2 + GAP, 2.1, cw2, 2.55, "The practice", [
    "One small example you can follow with a calculator",
    "Every formula by hand on the slide",
    "The same steps by hand in Python",
    "Check against polars' built-in functions"], "LuCode", "Practice card", head_h=0.45)
callout(s, 4.8, 0.6, "LuLightbulb", ("Why now: ", "regression output is made of these pieces: means, standard errors, "
                                                 "z-values and correlations."), "Why now", dark=True, size=12)

# 3  Running example
s = add_slide("Titel und Inhalt", "Our example: a chocolate kiosk",
              notes="A kiosk at the university sells one chocolate bar. Over five weeks the owner tried five prices. Small "
                    "numbers on purpose: every result today can be checked with a calculator. The daily sales of one week "
                    "(right) serve for mean and median.")
text(s, 0.5, 1.35, 4.4, 0.38, [[("Five weeks, five prices", {"bold": True, "col": NAVY, "size": HEAD, "head": True})]],
     "Weeks label", anchor=MSO_ANCHOR.MIDDLE)
grid(s, [("Week", "Price (EUR)", "Bars sold"), ("1", "2.00", "110"), ("2", "2.50", "100"), ("3", "3.00", "80"),
         ("4", "3.50", "70"), ("5", "4.00", "40")], [1.2, 1.6, 1.6], 0.5, 1.8, "Weeks", row_h=0.4, size=12.5,
     align=[L_, R, R])
text(s, 5.1, 1.35, 4.4, 0.38, [[("One week, day by day", {"bold": True, "col": NAVY, "size": HEAD, "head": True})]],
     "Days label", anchor=MSO_ANCHOR.MIDDLE)
grid(s, [("Day", "Bars sold"), ("Mon", "10"), ("Tue", "12"), ("Wed", "11"), ("Thu", "13"), ("Fri (event)", "54")],
     [2.2, 2.2], 5.1, 1.8, "Days", row_h=0.4, size=12.5, align=[L_, R])
callout(s, 4.45, 0.75, "LuCalculator", ("Small on purpose: ", "every number today can be checked with a calculator, "
                                                             "then in Python."), "Kiosk calc", dark=True)

# ---------------------------------------------------------------- Part 1 -----
divider("Scale types", "Part 1 of 5" + NB + "·" + NB + "nominal, ordinal, interval, ratio",
        "The scale decides which statistics and which charts make sense.")

# 5  Four scale types
s = add_slide("Titel und Tabelle", "Four scale types",
              notes="Nominal: categories without order (brand, country, payment method): count them, report the mode. "
                    "Ordinal: ordered categories with unequal or unknown distances (star rating, size S/M/L): median and "
                    "percentiles. Interval: equal distances but no true zero (temperature in °C): differences and means "
                    "are fine, ratios are not. Ratio: true zero (price, bars sold, revenue): everything, including 'twice "
                    "as much'. Interval and ratio together are called metric; nominal and ordinal categorical.")
table(s, [("Scale", "What you can say", "Kiosk example", "Typical statistics"),
          ("Nominal", "same or different", "brand, payment method", "counts, mode"),
          ("Ordinal", "more or less", "rating 1 to 5 stars, size S/M/L", "median, percentiles"),
          ("Interval", "how much more (no true zero)", "temperature in °C", "mean, standard deviation"),
          ("Ratio", "how many times (true zero)", "price, bars sold, revenue", "all, incl. ratios")],
      [1.5, 2.6, 2.8, 2.1], "Scales table", size=12.5, top=1.35, row_margin=60000)
callout(s, 4.0, 0.7, "LuLayers", ("Two groups: ", "nominal and ordinal are categorical; interval and ratio are metric "
                                                 "(numbers you can average)."), "Scale groups", dark=True)
text(s, 0.5, 4.85, 9.0, 0.4, [[("In regression: ", {"bold": True, "col": NAVY, "size": 12}),
                               ("categorical variables enter as dummies (0/1), metric variables as they are.",
                                {"size": 12})]], "Scale reg", anchor=MSO_ANCHOR.MIDDLE)

# 6  Why it matters + quiz
s = add_slide("Titel und Inhalt", "Why the scale matters",
              notes="20 °C is not twice as warm as 10 °C (in Kelvin it is 293 versus 283). An average of star ratings "
                    "treats an ordinal scale as metric: common in practice, but a decision you should state. An 'average "
                    "brand' means nothing. Quiz answers: country nominal; satisfaction 1 to 7 ordinal (often treated as "
                    "metric); age in years ratio; store ID nominal (a number, but only a label); discount yes/no nominal "
                    "(binary, a dummy); temperature interval.")
rows_with_bar(s, 0.5, 1.35, 4.4, [("20 °C ", "is not twice as warm as 10 °C: no true zero"),
                                  ("Mean of star ratings: ", "treats ordinal as metric; say so"),
                                  ("Store ID 17: ", "a number, but only a label: nominal")], "Matter", row_h=0.78)
box(s, 5.05, 1.35, 4.45, 0.45, ACC, "Quiz head")
text(s, 5.2, 1.35, 4.15, 0.45, [[("Quick quiz: which scale?", {"bold": True, "col": WHITE, "size": HEAD, "head": True})]],
     "Quiz head text", anchor=MSO_ANCHOR.MIDDLE)
box(s, 5.05, 1.8, 4.45, 2.06, LIGHT, "Quiz body")
text(s, 5.25, 1.9, 4.1, 1.9, ["Country of the customer", "Satisfaction from 1 to 7", "Age in years", "Discount: yes or no",
                              "Temperature on the day"], "Quiz items", bullets=True, size=12, space=5)
callout(s, 4.05, 0.7, "LuTarget", ("Rule: ", "check the scale before you average, chart or put a variable into a "
                                             "model."), "Matter rule", dark=True)

# ---------------------------------------------------------------- Part 2 -----
divider("Centre and spread", "Part 2 of 5" + NB + "·" + NB + "mean, median, variance, standard deviation, standard error",
        "Where are the data, how far do they scatter, and how precise is our average?")

# 8  Mean and median
s = add_slide("Titel und Inhalt", "Mean and median",
              notes="The mean adds all values and divides by n. The median is the middle value of the sorted data (with an "
                    "even n, the average of the two middle values). One event day with 54 bars pulls the mean to 20, but "
                    "the median stays at 12, which describes a normal day much better. Use the median for skewed data "
                    "(income, basket values, sales with promotions) and for ordinal scales; the mean for symmetric metric "
                    "data, and whenever totals matter (20 x 5 = 100 bars sold).")
code_block(s, 0.5, 1.35, 4.6, 0.95, "mean    x̄ = Σ x / n = 100 / 5 = 20\nmedian  middle of 10 11 12 13 54 = 12",
           "Mean formula", size=12)
rows_with_bar(s, 0.5, 2.45, 4.6, [("Mean ", "uses every value; one event day pulls it up"),
                                  ("Median ", "is robust: half the days below, half above"),
                                  ("Use the median ", "for skewed data and ordinal scales")], "MM rows", row_h=0.62)
picture(s, "ref_mean_median.png", 5.25, 1.35, 4.25, 2.9,
        "Dot chart of five days: four between 10 and 13 bars, Friday at 54; a solid line at the mean of 20 and a dashed "
        "line at the median of 12")
callout(s, 4.75, 0.55, "LuLightbulb", ("Totals need the mean: ", "20 bars a day × 5 days = 100 bars in the week."),
        "MM totals", size=12)

# 9  Variance and SD by hand
s = add_slide("Titel und Inhalt", "Variance and standard deviation, by hand",
              notes="Step by step for weekly bars sold: subtract the mean (80), square the deviations so that plus and minus "
                    "do not cancel, add them up (3,000), divide by n - 1 = 4: the variance is 750 (bars squared). The square "
                    "root brings us back to bars: the standard deviation is 27.4 bars, the typical distance of a week from "
                    "the mean. Why n - 1: the deviations are measured from the sample mean, which sits closer to the data "
                    "than the true mean; dividing by n - 1 corrects this. polars, pandas and statsmodels use n - 1 for "
                    "samples.")
grid(s, [("Week", "Bars x", "x − x̄", "(x − x̄)²"), ("1", "110", "30", "900"), ("2", "100", "20", "400"),
         ("3", "80", "0", "0"), ("4", "70", "−10", "100"), ("5", "40", "−40", "1,600"), ("Sum", "400", "0", "3,000")],
     [1.0, 1.15, 1.15, 1.35], 0.5, 1.35, "Var table", row_h=0.38, size=12, last_bold=True, align=[L_, R, R, R])
code_block(s, 5.4, 1.35, 4.1, 1.35, "x̄  = 400 / 5 = 80\ns² = 3,000 / (5 − 1) = 750\ns  = √750 = 27.4 bars",
           "Var formula", size=12, caption="Variance and standard deviation")
rows_with_bar(s, 5.4, 2.85, 4.1, [("Variance ", "is in squared units (bars²)"),
                                  ("Standard deviation ", "is in bars: the typical distance from the mean")], "Var rows",
              row_h=0.58)
callout(s, 4.35, 0.85, "LuSigma", ("Why n − 1? ", "deviations from the sample mean are a little too small; dividing by n − 1 "
                                                 "corrects this. polars and pandas do it by default."), "Var n1",
        dark=True, size=12)

# 10  What SD tells you
s = add_slide("Titel und Inhalt", "Same mean, different spread",
              notes="Two kiosks sell 80 bars a week on average. Kiosk A varies between 75 and 85 (sd 3.8), kiosk B between "
                    "40 and 120 (sd 31.6). The mean alone hides this. For planning: kiosk B needs far more safety stock, and "
                    "a forecast for kiosk B is much less certain.")
picture(s, "ref_spread.png", 0.5, 1.35, 5.0, 2.73,
        "Two rows of five dots: kiosk A between 75 and 85, kiosk B between 40 and 120, a dashed line at the common mean of 80")
tw = 3.85
for i, (fig, lab) in enumerate([("3.8", "bars: sd of kiosk A"), ("31.6", "bars: sd of kiosk B")]):
    y = 1.35 + i * (1.0 + 0.1)
    box(s, 5.65, y, tw, 1.0, LIGHT, f"SD stat box {i + 1}")
    stat(s, 5.65, y, tw, 1.0, fig, lab, f"SD stat {i + 1}", fig_w=1.3, size=12)
callout(s, 4.3, 0.8, "LuShoppingCart", ("For the manager: ", "both sell 80 a week, but kiosk B needs much more safety stock "
                                                             "and its forecasts are less certain."), "SD manager",
        dark=True, size=12)

# 11  Standard error
s = add_slide("Titel und Inhalt", "Standard error: how precise is the mean?",
              notes="The standard deviation describes the customers; the standard error describes the precision of the "
                    "average we calculate from a sample. SE = s / square root of n. Spending per customer has sd EUR 4: with "
                    "16 customers the SE of the mean is EUR 1, with 64 customers EUR 0.50. Four times the sample halves the "
                    "standard error. A 95 % confidence interval is roughly mean plus or minus 2 SE (1.96 exactly): EUR 12 "
                    "plus or minus 1.96 gives 10.04 to 13.96. Every 'std err' in a regression table is a standard error.")
code_block(s, 0.5, 1.35, 9.0, 0.6, "SE = s / √n          95% confidence interval ≈ x̄ ± 1.96 · SE", "SE formula", size=12)
tw3 = (9.0 - GAP * 2) / 3
se = [("EUR" + NB + "1.00", "SE with 16 customers: 4 / √16"), ("EUR" + NB + "0.50", "SE with 64 customers: 4 / √64"),
      ("10.04–13.96", "95" + PCT + " interval for a mean spend of EUR 12 (n = 16)")]
for i, (fig, lab) in enumerate(se):
    x = 0.5 + i * (tw3 + GAP)
    box(s, x, 2.1, tw3, 1.25, LIGHT, f"SE box {i + 1}")
    box(s, x, 2.1, 0.07, 1.25, ACC, f"SE bar {i + 1}")
    text(s, x + 0.2, 2.15, tw3 - 0.3, 0.55, [[(fig, {"size": 24, "bold": True, "col": NAVY, "head": True})]],
         f"SE fig {i + 1}", anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 2.7, tw3 - 0.3, 0.6, [lab], f"SE lab {i + 1}", size=12)
rows_with_bar(s, 0.5, 3.5, 9.0, [("sd ", "describes how much customers differ: spending varies by about EUR 4"),
                                 ("SE ", "describes how precise the average is: it shrinks with √n; four times the sample, "
                                         "half the SE")], "SE rows", row_h=0.5, gap=0.07)
callout(s, 4.7, 0.55, "LuTable", ("In every regression table: ", "the column 'std err' is the standard error of a "
                                                                "coefficient."), "SE reg", size=12)

# ---------------------------------------------------------------- Part 3 -----
divider("Standardisation and the normal distribution", "Part 3 of 5" + NB + "·" + NB + "z-scores, normal and standard "
        "normal distribution", "z-scores make different scales comparable; the normal distribution tells us how unusual a "
        "value is.")

# 13  Standardisation
s = add_slide("Titel und Inhalt", "Standardisation: z-scores",
              notes="A z-score says how many standard deviations a value lies above or below the mean. It removes the units, "
                    "so values on different scales become comparable. Vienna sold 140 bars in a week (mean 120, sd 10): "
                    "z = 2.0. Graz sold 75 (mean 60, sd 15): z = 1.0. Vienna's week was the more unusual one, although "
                    "Graz's increase looks large. In our kiosk, week 5 with 40 bars has z = (40 - 80) / 27.4 = -1.46. A "
                    "standardised variable always has mean 0 and standard deviation 1; standardised regression coefficients "
                    "use exactly this.")
code_block(s, 0.5, 1.35, 9.0, 0.6, "z = (x − x̄) / s        a standardised variable has mean 0 and standard deviation 1",
           "Z formula", size=12)
grid(s, [("Kiosk", "Bars this week", "Mean", "sd", "z"), ("Vienna", "140", "120", "10", "+2.0"),
         ("Graz", "75", "60", "15", "+1.0"), ("Our kiosk, week 5", "40", "80", "27.4", "−1.46")],
     [2.6, 1.7, 1.4, 1.4, 1.9], 0.5, 2.1, "Z table", row_h=0.45, size=12.5, align=[L_, R, R, R, R])
callout(s, 4.1, 0.75, "LuScale", ("Read it: ", "Vienna's week lies 2 standard deviations above normal, Graz's only 1: "
                                               "Vienna's week was the more unusual one."), "Z read", dark=True)
text(s, 0.5, 4.95, 9.0, 0.4, [[("Use: ", {"bold": True, "col": NAVY, "size": 12}),
                               ("compare across stores, countries or units; spot outliers (|z| > 3); standardised "
                                "coefficients.", {"size": 12})]], "Z use", anchor=MSO_ANCHOR.MIDDLE)

# 14  Normal distribution
s = add_slide("Titel und Inhalt", "The normal distribution",
              notes="A machine fills 100 g bars; the weights are normally distributed with mean 100 g and sd 2 g. The bell "
                    "curve is fully described by mean and standard deviation. The 68-95-99.7 rule: 68 % of bars lie within "
                    "one sd (98 to 102 g), 95 % within two (96 to 104 g; exactly 1.96 sd), 99.7 % within three (94 to "
                    "106 g). Many sums and averages are approximately normal even when single values are not (central "
                    "limit theorem): this is why means and regression coefficients can be tested with the normal curve.")
picture(s, "ref_normal.png", 0.5, 1.35, 4.6, 3.0,
        "Bell curve of bar weights centred at 100 g; the band from 98 to 102 g shaded dark (68 %), from 96 to 104 g light "
        "(95 %)")
rule = [("68" + PCT, "within ±1 sd: 98 to 102 g"), ("95" + PCT, "within ±2 sd: 96 to 104 g"),
        ("99.7" + PCT, "within ±3 sd: 94 to 106 g")]
for i, (fig, lab) in enumerate(rule):
    y = 1.35 + i * (0.8 + 0.08)
    box(s, 5.25, y, 4.25, 0.8, LIGHT, f"Rule box {i + 1}")
    stat(s, 5.25, y, 4.25, 0.8, fig, lab, f"Rule stat {i + 1}", fig_w=1.8, size=12)
callout(s, 4.5, 0.75, "LuLightbulb", ("Why it matters: ", "averages and regression coefficients are approximately normal "
                                                         "(central limit theorem), so we can test them."), "Normal why",
        dark=True, size=12)

# 15  Standard normal
s = add_slide("Titel und Inhalt", "The standard normal distribution",
              notes="Standardising any normal variable gives the standard normal distribution: mean 0, sd 1. One table (or "
                    "one line of Python) then answers every question. Share of bars below 97 g: z = (97 - 100) / 2 = -1.5, "
                    "P(Z < -1.5) = 6.7 %. The value 1.96 cuts off 2.5 % in each tail: it is the 1.96 in confidence "
                    "intervals and the critical value of a two-sided test at 5 %.")
picture(s, "ref_std_normal.png", 0.5, 1.35, 4.6, 3.0,
        "Standard normal curve from -3.5 to 3.5 with both tails beyond 1.96 shaded, each labelled 2.5 %, the centre 95 %")
grid(s, [("z", "share below z"), ("1.00", "84.1" + PCT), ("1.645", "95.0" + PCT), ("1.96", "97.5" + PCT),
         ("2.576", "99.5" + PCT)], [1.6, 2.6], 5.25, 1.35, "Z share", row_h=0.4, size=12.5, align=[L_, R])
box(s, 5.25, 3.5, 4.25, 0.85, LIGHT, "Z worked box")
box(s, 5.25, 3.5, 0.07, 0.85, ACC, "Z worked bar")
text(s, 5.45, 3.5, 4.0, 0.85, [("Bars under 97 g? ", "z = (97 − 100) / 2 = −1.5, so 6.7" + PCT + " of all bars")],
     "Z worked", anchor=MSO_ANCHOR.MIDDLE, size=12)
callout(s, 4.55, 0.7, "LuTarget", ("1.96: ", "the number behind 95" + PCT + " confidence intervals and two-sided tests at "
                                            "5" + PCT + "."), "Z196", dark=True, size=12)

# ---------------------------------------------------------------- Part 4 -----
divider("Covariance and correlation", "Part 4 of 5" + NB + "·" + NB + "do two variables move together, and how strongly?",
        "From one variable to two: price and bars sold at the kiosk.")

# 17  Covariance
s = add_slide("Titel und Inhalt", "Covariance: do price and sales move together?",
              notes="Multiply the deviation of price by the deviation of bars sold for every week. When one is above its "
                    "mean while the other is below, the product is negative (weeks 1, 2, 4 and 5); week 3 sits exactly on "
                    "both means. The sum is -85; divided by n - 1 = 4 the covariance is -21.25. The sign gives the "
                    "direction: negative, higher prices go with fewer bars. The size depends on the units (EUR x bars), so "
                    "-21.25 alone says little about strength.")
picture(s, "ref_quadrants.png", 0.5, 1.35, 4.4, 3.06,
        "Scatter plot of the five weeks, price against bars sold, split by dashed lines at the means (3.00 EUR and 80 "
        "bars); the points lie in the top-left and bottom-right quadrants, where the product of deviations is negative")
grid(s, [("Week", "p − p̄", "x − x̄", "product"), ("1", "−1.0", "30", "−30"), ("2", "−0.5", "20", "−10"),
         ("3", "0", "0", "0"), ("4", "0.5", "−10", "−5"), ("5", "1.0", "−40", "−40"), ("Sum", "", "", "−85")],
     [0.95, 1.1, 1.1, 1.3], 5.05, 1.35, "Cov table", row_h=0.36, size=12, last_bold=True, align=[L_, R, R, R])
code_block(s, 5.05, 4.05, 4.45, 0.55, "cov = −85 / (5 − 1) = −21.25", "Cov formula", size=12)
callout(s, 4.75, 0.55, "LuShuffle", ("Sign = direction; ", "the size depends on the units (EUR × bars)."), "Cov sign",
        size=12)

# 18  Correlation
s = add_slide("Titel und Inhalt", "Correlation: covariance without units",
              notes="Dividing the covariance by both standard deviations removes the units: r = -21.25 / (0.79 x 27.39) = "
                    "-0.98. r always lies between -1 and +1; -0.98 is an extremely strong negative relationship (rule of "
                    "thumb in Bojinov et al.: above 0.8 extremely strong). The gallery shows strong positive, none, strong "
                    "negative, and a U-shape where r is zero although x and y are clearly related: r only measures linear "
                    "relationships.")
code_block(s, 0.5, 1.35, 9.0, 0.6, "r = cov / (s_p · s_x) = −21.25 / (0.79 · 27.39) = −0.98          −1 ≤ r ≤ +1",
           "R formula", size=12)
picture(s, "ref_corr_gallery.png", 0.5, 2.05, 9.0, 2.3,
        "Four scatter plots: a rising cloud with r about 0.9, a shapeless cloud with r about 0, a falling cloud with r about "
        "-0.9, and a U-shaped curve with r about 0")
callout(s, 4.5, 0.7, "LuTriangleAlert", ("Two warnings: ", "r only measures straight-line relationships (the U has r ≈ 0), "
                                                          "and correlation is not causation."), "R warn", dark=True,
        size=12)
source(s, "Bojinov, Parzen & Hamilton 2025 (HBS note 9-622-100)")

# 19  Bridge to regression
s = add_slide("Titel und Inhalt", "From correlation to regression",
              notes="Ice-cream sales and sunburn correlate because sunshine drives both: correlation is not causation. The "
                    "regression slope through the five weeks is cov / var(price) = -21.25 / 0.625 = -34: each euro more, "
                    "34 bars fewer. The same building blocks (means, deviations, covariance, variance) produce the "
                    "least-squares line of the next deck.")
rows_with_bar(s, 0.5, 1.35, 9.0, [("Not causation: ", "ice-cream sales and sunburn rise together; sunshine drives both"),
                                  ("Correlation ", "says how strongly two variables move together, from −1 to +1"),
                                  ("Regression ", "says by how much y changes when x changes by one unit")],
              "Bridge rows", row_h=0.6, gap=0.08)
code_block(s, 0.5, 3.55, 9.0, 0.75, "slope b = cov(p, x) / var(p) = −21.25 / 0.625 = −34 bars per EUR\n"
                                    "line through the means: x = 80 − 34 · (p − 3)", "Slope formula", size=12)
callout(s, 4.5, 0.7, "LuChartLine", ("Next: ", "the least-squares line, built from exactly these pieces."), "Bridge next",
        dark=True)

# ---------------------------------------------------------------- Part 5 -----
divider("Python in Positron", "Part 5 of 5" + NB + "·" + NB + "mean, standard deviation, z-scores, covariance and "
        "correlation by hand, then checked", "Open sessions/01-foundations/stats_refresher.qmd and run it cell by cell "
        "with Ctrl/Cmd + Enter. Each step first calculates by hand, then checks with the built-in function.")

# 21  Step 1 mean
s = add_slide("Titel und Inhalt", "Step 1: the data and the mean",
              notes="A polars data frame from three lists. The mean by hand is the sum divided by the number of rows; "
                    "polars' mean() gives the same 80.0.")
code_block(s, 0.5, 1.35, 5.4, 3.0, """import polars as pl

kiosk = pl.DataFrame({
    "week":  [1, 2, 3, 4, 5],
    "price": [2.0, 2.5, 3.0, 3.5, 4.0],
    "units": [110, 100, 80, 70, 40],
})

n = kiosk.height                         # 5
mean_units = kiosk["units"].sum() / n    # by hand
kiosk["units"].mean()                    # check
kiosk["units"].median()""", "Mean code", size=11, caption="Python: polars")
grid(s, [("Result", "Value"), ("n", "5"), ("mean, by hand", "80.0"), ("mean(), check", "80.0"), ("median()", "80.0")],
     [2.0, 1.45], 6.05, 1.35, "Mean out", row_h=0.4, size=12, align=[L_, R])
callout(s, 4.55, 0.6, "LuPlay", ("Run it: ", "click into a cell and press Ctrl/Cmd + Enter; the result appears in the "
                                             "Console."), "Run it", size=12)

# 22  Step 2 sd
s = add_slide("Titel und Inhalt", "Step 2: variance, standard deviation, standard error",
              notes="The same steps as on the slide by hand: deviations, squares, sum, divide by n - 1, square root. polars' "
                    "std() uses n - 1 by default (ddof=1), so it matches. The standard error of the mean is sd / sqrt(n).")
code_block(s, 0.5, 1.35, 5.4, 3.0, """dev = kiosk["units"] - mean_units
dev        # 30, 20, 0, -10, -40

var_units = (dev ** 2).sum() / (n - 1)   # 750.0
sd_units = var_units ** 0.5              # 27.39
kiosk["units"].std()                     # check: n - 1

se_units = sd_units / n ** 0.5           # 12.25""", "SD code", size=11, caption="Python: by hand, then checked")
grid(s, [("Result", "Value"), ("variance", "750.0"), ("sd, by hand", "27.39"), ("std(), check", "27.39"),
         ("standard error", "12.25")], [2.0, 1.45], 6.05, 1.35, "SD out", row_h=0.4, size=12, align=[L_, R])
callout(s, 4.55, 0.6, "LuSigma", ("Note: ", "std() divides by n − 1 (ddof=1), like the formula on the slide."), "SD note",
        size=12)

# 23  Step 3 z
s = add_slide("Titel und Inhalt", "Step 3: standardise",
              notes="with_columns adds a new column; the expression subtracts the mean and divides by the standard deviation. "
                    "The z-scores have mean 0 and standard deviation 1. Week 5 lies 1.46 standard deviations below the mean.")
code_block(s, 0.5, 1.35, 5.4, 2.7, """kiosk = kiosk.with_columns(
    z_units=(pl.col("units") - pl.col("units").mean())
            / pl.col("units").std()
)
kiosk

kiosk["z_units"].mean()    # 0
kiosk["z_units"].std()     # 1""", "Z code", size=11, caption="Python: z-scores")
grid(s, [("Week", "Bars", "z"), ("1", "110", "1.10"), ("2", "100", "0.73"), ("3", "80", "0.00"), ("4", "70", "−0.37"),
         ("5", "40", "−1.46")], [1.0, 1.15, 1.3], 6.05, 1.35, "Z out", row_h=0.4, size=12, align=[L_, R, R])
callout(s, 4.55, 0.6, "LuScale", ("Read it: ", "week 5 lies 1.46 standard deviations below an average week."), "Z out read",
        size=12)

# 24  Step 4 cov and r
s = add_slide("Titel und Inhalt", "Step 4: covariance and correlation",
              notes="Deviations of both variables, multiplied week by week, summed and divided by n - 1: the covariance. "
                    "Divided by both standard deviations: the correlation. pl.cov and pl.corr check both in one line.")
code_block(s, 0.5, 1.35, 5.4, 3.0, """dev_p = kiosk["price"] - kiosk["price"].mean()
dev_u = kiosk["units"] - kiosk["units"].mean()

cov = (dev_p * dev_u).sum() / (n - 1)          # -21.25
r = cov / (kiosk["price"].std()
           * kiosk["units"].std())             # -0.981

kiosk.select(
    pl.cov("price", "units").alias("cov"),
    pl.corr("price", "units").alias("r"))      # check""", "Cov code", size=11, caption="Python: by hand, then checked")
grid(s, [("Result", "Value"), ("covariance", "−21.25"), ("correlation", "−0.981"), ("pl.cov, check", "−21.25"),
         ("pl.corr, check", "−0.981")], [2.0, 1.45], 6.05, 1.35, "Cov out", row_h=0.4, size=12, align=[L_, R])
callout(s, 4.55, 0.6, "LuShuffle", ("Same numbers ", "as by hand: the computer only saves the typing."), "Cov same",
        size=12)

# 25  Step 5 chart
s = add_slide("Titel und Inhalt", "Step 5: look at it",
              notes="plotnine takes the polars data frame directly. Always plot two variables before you trust a correlation: "
                    "the U-shape on the correlation slide had r = 0.")
code_block(s, 0.5, 1.35, 4.6, 2.2, """from plotnine import (ggplot, aes,
                      geom_point, labs)

(ggplot(kiosk, aes("price", "units"))
 + geom_point(size=3)
 + labs(x="Price (EUR)",
        y="Bars sold per week"))""", "Plot code", size=11, caption="Python: plotnine")
picture(s, "ref_quadrants.png", 5.25, 1.35, 4.25, 2.96, "The scatter plot of price against bars sold for the five weeks")
callout(s, 4.5, 0.75, "LuEye", ("Plot first: ", "a correlation can hide a curve, an outlier or two groups."), "Plot first",
        dark=True)

# 26  Your turn
s = add_slide("Titel und Inhalt", "Your turn: by hand, then in Python",
              notes="Answers: 1 day nominal-ordinal (weekday order), temperature interval, bars sold ratio. 2 mean 48, "
                    "median 50. 3 deviations 12, 7, 2, -8, -13; squares sum 430; variance 107.5; sd 10.37. 4 z of the "
                    "hottest day: (35 - 48) / 10.37 = -1.25. 5 temperature deviations -10, -5, 0, 5, 10; products -120, -35, "
                    "0, -40, -130, sum -325; covariance -81.25; sd temperature 7.91; r = -0.99. Warmer days, fewer "
                    "chocolate bars.")
stepper(s, 0.5, 1.35, 9.0, 0.5, ["Pairs, paper:\v10 minutes", "Python check:\v10 minutes"], "Turn steps", size=12)
grid(s, [("Day", "Temp. °C", "Bars"), ("Mon", "10", "60"), ("Tue", "15", "55"), ("Wed", "20", "50"), ("Thu", "25", "40"),
         ("Fri", "30", "35")], [0.95, 1.2, 1.0], 0.5, 2.0, "Turn data", row_h=0.4, size=12, align=[L_, R, R])
numbered(s, 3.85, 2.0, 5.65, [
    ("Scales: ", "which scale is each column?"),
    ("Centre: ", "mean and median of bars sold"),
    ("Spread: ", "variance and sd of bars, by hand"),
    ("z: ", "z-score of the hottest day's sales"),
    ("Together: ", "covariance and r of temperature and bars; then check all in Python")],
    "Turn q", row_h=0.5, gap=0.07, size=12)

# 27  Takeaways
s = add_slide("Titel und Inhalt", "Key takeaways",
              notes="The vocabulary for the rest of the course: every regression table reports means, standard errors, "
                    "z- or t-values and builds on covariances.")
numbered(s, 0.5, 1.35, 9.0, [
    ("Scale ", "first: it decides which statistics make sense"),
    ("Mean and median: ", "the median resists outliers; totals need the mean"),
    ("sd ", "describes the data; SE = sd / √n describes the precision of the mean"),
    ("z = (x − x̄) / s ", "makes values comparable; 95" + PCT + " of a normal lies within ±1.96 sd"),
    ("Covariance ", "gives the direction, correlation the strength (−1 to +1), regression the size")],
    "Takeaways", row_h=0.64, gap=0.09, size=12)
callout(s, 5.0, 0.5, "LuChartLine", ("Next: ", "linear regression on real chocolate data."), "Take next", dark=True, size=12)

# 28  References
s = add_slide("Titel und Inhalt", "References",
              notes="Field is the friendliest textbook for students without a statistics background; Seeing Theory and "
                    "StatQuest visualise every concept of today.")
refs_l = [("Field, A. (2024). ", "Discovering Statistics Using IBM SPSS Statistics, 6th ed. SAGE. Chapters 1 to 2 and 8."),
          ("Bojinov, I. I., Parzen, M. & Hamilton, P. J. (2025). ", "Linear Regression. HBS note 9-622-100: correlation."),
          ("Hair, J. F. et al. (2019). ", "Multivariate Data Analysis, 8th ed. Cengage. Chapter 2."),
          ("Stock, J. H. & Watson, M. W. (2019). ", "Introduction to Econometrics, 4th ed. Pearson. Chapters 2 to 3.")]
refs_r = [("Seeing Theory. ", "Brown University: interactive visual introduction to probability and statistics (free)."),
          ("StatQuest. ", "Videos on the normal distribution, standard deviation vs standard error, covariance and "
                          "correlation (free)."),
          ("polars documentation. ", "mean, median, std, cov, corr."),
          ("Gelman, Hill & Vehtari (2020). ", "Regression and Other Stories, ch. 3 to 4 (free PDF).")]
cw2 = (9.0 - GAP) / 2
for k, (head, ic, refs) in enumerate([("Books and readings", "LuBookOpen", refs_l), ("Online", "LuCode", refs_r)]):
    x = 0.5 + k * (cw2 + GAP)
    box(s, x, 1.35, cw2, 0.45, ACC, f"Ref {k + 1} header")
    text(s, x + 0.15, 1.35, cw2 - 0.75, 0.45, [head], f"Ref {k + 1} heading", size=HEAD, col=WHITE, bold=True, head=True,
         anchor=MSO_ANCHOR.MIDDLE)
    icon(s, ic, "white", x + cw2 - 0.48, 1.4, 0.34, f"ref {k + 1}")
    box(s, x, 1.8, cw2, 3.75, LIGHT, f"Ref {k + 1} body")
    text(s, x + 0.18, 1.9, cw2 - 0.33, 3.6, [[(lead, {"bold": True, "col": NAVY}), (rest, {})] for lead, rest in refs],
         f"Ref {k + 1} text", size=11.5, space=6)

# 29  Closing
add_slide("Abschlussfolie Kontakt", notes="Questions? Contact details on the card; office hours via the booking link.")

deck.save(OUT)
