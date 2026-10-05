"""Shared building blocks for the course decks on the WU template (polish-slides skill).

Design follows .claude/skills/polish-slides/references/design-patterns.md: icon rows,
cards with an accent1 header band, steppers, stat rows, all in theme colours.
Type sizes: 13 pt text, 15 pt card and table headings, 16 pt slide titles, 28 pt figures.

    deck = Deck(TEMPLATE, FOOTER)
    s = deck.add_slide("Titel und Inhalt", "Title", notes="...")
    box(s, ...); text(s, ...); card(s, ...)
    deck.save(OUT)
"""
import copy
import io
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

ICONS = Path(__file__).resolve().parent / "icons"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NB = " "
BODY, HEAD, TITLE, FIG = 13, 15, 16, 28
ACC, NAVY, LIGHT, INK, WHITE, BARS = (MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.TEXT_2, MSO_THEME_COLOR.BACKGROUND_2,
                                      MSO_THEME_COLOR.TEXT_1, MSO_THEME_COLOR.BACKGROUND_1, MSO_THEME_COLOR.ACCENT_4)
GAP = 0.15
MONO = "Consolas"


def I(x):
    return Emu(int(round(x * 914400)))


def set_runs(p, text):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if isinstance(text, tuple):
        lead, rest = text
        r1 = p.add_run(); r1.text = lead; r1.font.bold = True
        r2 = p.add_run(); r2.text = rest
    else:
        p.add_run().text = text


def ph(slide, idx):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == idx:
            return sh
    raise KeyError(idx)


def remove_shape(sh):
    sh._element.getparent().remove(sh._element)


class Deck:
    """A presentation started from the template with its sample slides removed."""

    def __init__(self, template, footer):
        self.prs = Presentation(template)
        self.footer = footer
        for sldId in list(self.prs.slides._sldIdLst):
            self.prs.part.drop_rel(sldId.rId)
            self.prs.slides._sldIdLst.remove(sldId)
        self.L = {l.name: l for l in self.prs.slide_layouts}

    def add_slide(self, layout_name, title=None, footer="default", notes=None, keep_body=False):
        footer = self.footer if footer == "default" else footer
        s = self.prs.slides.add_slide(self.L[layout_name])
        for lph in self.L[layout_name].placeholders:   # python-pptx does not copy footer and slide number
            if lph.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                s.shapes._spTree.append(copy.deepcopy(lph._element))
        for sh in list(s.placeholders):
            t = sh.placeholder_format.type
            if t == PP_PLACEHOLDER.FOOTER and footer is not None:
                set_runs(sh.text_frame.paragraphs[0], footer)
            elif t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and title is not None:
                set_runs(sh.text_frame.paragraphs[0], title)
                for r in sh.text_frame.paragraphs[0].runs:
                    r.font.size = Pt(TITLE)
            elif t == PP_PLACEHOLDER.OBJECT and not keep_body and layout_name == "Titel und Inhalt":
                remove_shape(sh)                      # designed slides draw their own content
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, out):
        """Tag every run as British English, then save."""
        for slide in self.prs.slides:
            for sh in slide.shapes:
                frames = [sh.text_frame] if sh.has_text_frame else []
                if sh.has_table:
                    frames += [c.text_frame for row in sh.table.rows for c in row.cells]
                for tf in frames:
                    for p in tf.paragraphs:
                        for r in p.runs:
                            r._r.get_or_add_rPr().set("lang", "en-GB")
        self.prs.save(out)
        print("saved", out, "slides:", len(self.prs.slides))


def color(fmt, c):
    if isinstance(c, str):
        fmt.rgb = RGBColor.from_string(c)
    else:
        fmt.theme_color = c


def box(slide, x, y, w, h, fill, name, shape=MSO_SHAPE.RECTANGLE, line=None):
    sh = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sh.name = name
    sh.fill.solid(); color(sh.fill.fore_color, fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.theme_color = line
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def text(slide, x, y, w, h, paras, name, size=BODY, col=INK, bold=False, head=False, anchor=MSO_ANCHOR.TOP,
         align=PP_ALIGN.LEFT, bullets=False, space=4, lead_col=NAVY, margin=0.0):
    """paras: list of str, (lead, rest) tuples or lists of (text, {bold, col, size, head, font, link}) runs."""
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = I(margin)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        runs = para if isinstance(para, list) else ([(para[0], {"bold": True, "col": lead_col}), (para[1], {})]
                                                     if isinstance(para, tuple) else [(para, {})])
        for t, fmt in runs:
            r = p.add_run(); r.text = t
            r.font.size = Pt(fmt.get("size", size))
            r.font.bold = fmt.get("bold", bold)
            color(r.font.color, fmt.get("col", col))
            if head or fmt.get("head"):
                r.font.name = "+mj-lt"
            if fmt.get("font"):
                r.font.name = fmt["font"]
            if fmt.get("link"):
                r.hyperlink.address = fmt["link"]
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(I(0.2))); pPr.set("indent", str(-I(0.2)))
            clr = etree.SubElement(pPr, f"{{{A}}}buClr"); sc = etree.SubElement(clr, f"{{{A}}}schemeClr"); sc.set("val", "accent1")
            etree.SubElement(pPr, f"{{{A}}}buFont").set("typeface", "Arial")
            etree.SubElement(pPr, f"{{{A}}}buChar").set("char", "▪")
    return tb


def icon(slide, name, tone, x, y, size, label):
    pic = slide.shapes.add_picture(str(ICONS / f"{name}-{tone}.png"), I(x), I(y), I(size), I(size))
    pic.name = f"Icon {label}"
    pic._element.nvPicPr.cNvPr.set("descr", "")      # decorative
    return pic


def card(slide, x, y, w, h, heading, items, icon_name, name, head_h=0.5, stat=None, bullets=True):
    """Card with a solid accent1 header band, white icon at its right corner and a bg2 body."""
    box(slide, x, y + head_h, w, h - head_h, LIGHT, f"{name} body")
    box(slide, x, y, w, head_h, ACC, f"{name} header")
    text(slide, x + 0.15, y, w - 0.75, head_h, [heading], f"{name} heading", size=HEAD, col=WHITE, bold=True,
         head=True, anchor=MSO_ANCHOR.MIDDLE)
    icon(slide, icon_name, "white", x + w - 0.48, y + (head_h - 0.34) / 2, 0.34, name)
    body_h = h - head_h - 0.3 - (0.75 if stat else 0)
    text(slide, x + 0.18, y + head_h + 0.15, w - 0.33, body_h, items, f"{name} text", bullets=bullets, space=6)
    if stat:
        fig, label = stat
        text(slide, x + 0.18, y + h - 0.8, w - 0.33, 0.65,
             [[(fig, {"size": FIG, "bold": True, "col": NAVY, "head": True}), ("  " + label, {"col": INK})]],
             f"{name} stat", anchor=MSO_ANCHOR.BOTTOM)


def stepper(slide, x, y, w, h, steps, name, current=None, size=BODY):
    n = len(steps)
    overlap = 0.12
    sw = (w + overlap * (n - 1)) / n
    for i, label in enumerate(steps):
        shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        is_cur = current is not None and i == current
        sh = box(slide, x + i * (sw - overlap), y, sw, h, NAVY if is_cur else LIGHT, f"{name} step {i + 1}", shape=shape)
        sh.adjustments[0] = 0.28
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = I(0.26 if i else 0.1); tf.margin_right = I(0.14)
        tf.margin_top = tf.margin_bottom = I(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        p.text = label
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = True
            color(r.font.color, WHITE if is_cur else NAVY)


def code_block(slide, x, y, w, h, code, name, size=11, caption=None):
    """Navy box with monospaced white code; an optional small caption bar (file or terminal) on top."""
    top = y
    if caption:
        box(slide, x, y, w, 0.3, ACC, f"{name} caption bar")
        text(slide, x + 0.12, y, w - 0.24, 0.3, [[(caption, {"bold": True, "col": WHITE, "size": 11})]],
             f"{name} caption", anchor=MSO_ANCHOR.MIDDLE)
        top = y + 0.3
    box(slide, x, top, w, h - (top - y), NAVY, f"{name} box")
    lines = code.strip("\n").split("\n")
    text(slide, x + 0.15, top + 0.1, w - 0.3, h - (top - y) - 0.2,
         [[(ln if ln else " ", {"font": MONO, "col": WHITE})] for ln in lines], f"{name} code", size=size, space=0)


def copy_pictures(src_shape, dst_slide):
    el = copy.deepcopy(src_shape._element)
    for blip in el.iter(f"{{{A}}}blip"):
        rid = blip.get(f"{{{R}}}embed")
        if rid:
            _, new = dst_slide.part.get_or_add_image_part(io.BytesIO(src_shape.part.related_part(rid).blob))
            blip.set(f"{{{R}}}embed", new)
    dst_slide.shapes._spTree.append(el)
