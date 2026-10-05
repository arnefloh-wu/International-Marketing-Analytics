"""Apply the deck's type sizes to the instructor profile slides added by the
add-instructor-slide skill: slide titles 16 pt, card headings 15 pt.

usage: python slides/finish_profile_slides.py DECK.pptx
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt

path = sys.argv[1]
prs = Presentation(path)
done = 0
for slide in prs.slides:
    titles = [sh for sh in slide.placeholders
              if sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)]
    title = titles[0] if titles else None     # the asset's title placeholder has a non-zero idx
    if title is None or not title.text_frame.text.startswith("Module Convenor"):
        continue
    for p in title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(16)
    for sh in slide.shapes:
        if sh.name.startswith("Card header") and sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(15)
    done += 1
prs.save(path)
print(f"profile slides adjusted: {done}")
